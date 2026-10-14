"""Append validated checks to the bottom of each payee sheet, matching the sheet's own layout."""
import os, re, shutil, datetime, tempfile
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

HEADERS = {"customer": "CUSTOMER", "ticket": "TICKETREF", "payref": "PAYMENTREF", "taxable": "TAXABLE",
           "cash": "CASH", "check": "CHECK", "credit": "CREDIT", "depbank": "DEPOSITBANK"}
DEFAULT_COLS = {"customer": 1, "ticket": 2, "payref": 3, "taxable": 4, "cash": 5, "check": 6, "credit": 7}
DAYPFX = re.compile(r"^\s*(\d{1,2})/(\d{1,2})\b")

class WorkbookError(Exception): pass

def nh(v): return re.sub(r"[\s:.]", "", str(v or "")).upper()

def list_sheets(path, ignore):
    wb = openpyxl.load_workbook(path, read_only=True)
    try: names = list(wb.sheetnames)
    finally: wb.close()
    ig = {nh(x) for x in ignore}
    return [n for n in names if nh(n) not in ig and not nh(n).startswith("TAXE")]

def find_layout(ws):
    """Locate the last header row containing CUSTOMER and map column positions."""
    hdr_row, cols = None, {}
    for r in range(1, min(ws.max_row, 20000) + 1):
        for c in range(1, min(ws.max_column, 30) + 1):
            if nh(ws.cell(r, c).value) == "CUSTOMER":
                m = {}
                for cc in range(1, min(ws.max_column, 30) + 1):
                    for k, h in HEADERS.items():
                        if nh(ws.cell(r, cc).value) == h: m[k] = cc
                hdr_row, cols = r, m
                break
    if hdr_row is None: return 1, dict(DEFAULT_COLS)
    return hdr_row, {**DEFAULT_COLS, **cols}

def last_rows(ws, hdr, cols):
    """(last row with a customer value, last row that looks like a check row)."""
    last, last_check = hdr, None
    cc, kc = cols["customer"], cols["check"]
    for r in range(ws.max_row, hdr, -1):
        v = ws.cell(r, cc).value
        if v not in (None, ""):
            if last == hdr: last = r
            if DAYPFX.match(str(v)) and isinstance(ws.cell(r, kc).value, (int, float)):
                last_check = r; break
    return last, last_check

def num_or_text(v):
    v = str(v or "").strip()
    return int(v) if v.isdigit() and not v.startswith("0") and len(v) < 15 else (v or None)

def check_lock(path):
    d, n = os.path.split(path)
    if os.path.exists(os.path.join(d, "~$" + n)):
        raise WorkbookError("The workbook is open in Excel. Close it (or save and close) and press Send again.")

def ensure_depbank(ws, hdr, cols):
    """Find the DEPOSIT BANK column, or add one at the end of the header row."""
    if "depbank" in cols: return cols["depbank"]
    filled = [c for c in range(1, ws.max_column + 1) if ws.cell(hdr, c).value not in (None, "")]
    last = max(filled) if filled else max(cols.values())
    c = last + 1
    ws.cell(hdr, c).value = "DEPOSIT BANK"
    ws.cell(hdr, c)._style = copy(ws.cell(hdr, last)._style)
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 22
    cols["depbank"] = c
    return c

def day_of(it):
    """(year, month, day) the check was RECEIVED; falls back to the check date for old records."""
    r = it.get("received") or ""
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", r): return int(r[:4]), int(r[5:7]), int(r[8:10])
    d = it["date"]; return 2000 + int(d[6:8]), int(d[:2]), int(d[3:5])

def backup(path, keep=20, min_age=0):
    d = os.path.join(os.path.dirname(path), "checkflow_backups")
    os.makedirs(d, exist_ok=True)
    if min_age:
        import time
        recent = [os.path.getmtime(os.path.join(d, f)) for f in os.listdir(d)]
        if recent and time.time() - max(recent) < min_age: return None
    b = os.path.join(d, f"{os.path.splitext(os.path.basename(path))[0]}_{datetime.datetime.now():%Y%m%d_%H%M%S}{os.path.splitext(path)[1]}")
    shutil.copy2(path, b)
    old = sorted(os.listdir(d))
    for f in old[:-keep]:
        try: os.remove(os.path.join(d, f))
        except OSError: pass
    return b

def write_checks(path, items, day_gap=1):
    """items: [{id, sheet, date 'MM/DD/YY', payer, ticket, check, taxable, amount}] -> {id: (sheet,row)}"""
    if not path or not os.path.exists(path): raise WorkbookError("Workbook not found. Set the workbook path in Settings.")
    check_lock(path)
    keep_vba = path.lower().endswith(".xlsm")
    try:
        wb = openpyxl.load_workbook(path, keep_vba=keep_vba)
    except PermissionError:
        raise WorkbookError("The workbook is open in Excel. Close it and press Send again.")
    result = {}
    by_sheet = {}
    for it in items: by_sheet.setdefault(it["sheet"], []).append(it)
    for sheet, its in by_sheet.items():
        if sheet not in wb.sheetnames: raise WorkbookError(f"Sheet '{sheet}' does not exist in the workbook.")
        ws = wb[sheet]
        hdr, cols = find_layout(ws)
        last, tmpl = last_rows(ws, hdr, cols)
        row = last
        prev = DAYPFX.match(str(ws.cell(last, cols["customer"]).value or ""))
        prev_day = (int(prev.group(1)), int(prev.group(2))) if prev else None
        for it in sorted(its, key=lambda x: (day_of(x), x["id"])):
            _, mo, d = day_of(it)
            if prev_day != (mo, d): row += 1 + (day_gap if prev_day is not None else 0)
            else: row += 1
            prev_day = (mo, d)
            vals = {"customer": f"{mo}/{d} {it['payer'].upper()}", "ticket": num_or_text(it["ticket"]),
                    "payref": num_or_text(it["check"]), "taxable": "YES" if it["taxable"] == "Yes" else None, "check": float(it["amount"])}
            if tmpl:                                   # copy look of a real check row (fills, borders, fonts)
                for c in range(1, max(cols.values()) + 1):
                    ws.cell(row, c)._style = copy(ws.cell(tmpl, c)._style)
            for k, v in vals.items():
                ws.cell(row, cols[k]).value = v
            ws.cell(row, cols["check"]).number_format = "#,##0.00"
            if it.get("deposit_bank"):
                ws.cell(row, ensure_depbank(ws, hdr, cols)).value = it["deposit_bank"]
            result[it["id"]] = (sheet, row)
    backup(path)
    fd, tmp = tempfile.mkstemp(suffix=os.path.splitext(path)[1], dir=os.path.dirname(path))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)                          # atomic swap; fails if Excel holds the file
    except PermissionError:
        raise WorkbookError("The workbook is open in Excel. Close it and press Send again.")
    finally:
        if os.path.exists(tmp):
            try: os.remove(tmp)
            except OSError: pass
    return result

def update_rows(path, updates):
    """Update TAXABLE / DEPOSIT BANK on rows already written. updates: [{id,sheet,row,payer,check,taxable,deposit_bank}].
    Returns {id: ('ok', row) | ('err', message)}. The row is re-verified by customer + payment ref before touching it."""
    if not path or not os.path.exists(path): raise WorkbookError("Workbook not found. Set the workbook path in Settings.")
    check_lock(path)
    try: wb = openpyxl.load_workbook(path, keep_vba=path.lower().endswith(".xlsm"))
    except PermissionError: raise WorkbookError("The workbook is open in Excel.")
    out = {}
    for u in updates:
        if u["sheet"] not in wb.sheetnames: out[u["id"]] = ("err", f"Sheet '{u['sheet']}' not found."); continue
        ws = wb[u["sheet"]]; hdr, cols = find_layout(ws)
        def match(r):
            cust = str(ws.cell(r, cols["customer"]).value or "").upper()
            return cust.endswith(u["payer"].upper()) and str(ws.cell(r, cols["payref"]).value or "").strip() == str(u["check"]).strip()
        row = u["row"] if u["row"] and match(u["row"]) else next((r for r in range(ws.max_row, hdr, -1) if match(r)), None)
        if not row: out[u["id"]] = ("err", "Could not find this check's row in the sheet (was it moved or deleted?)."); continue
        ws.cell(row, cols["taxable"]).value = "YES" if u["taxable"] == "Yes" else None
        if u["deposit_bank"] or "depbank" in cols:
            ws.cell(row, ensure_depbank(ws, hdr, cols)).value = u["deposit_bank"] or None
        out[u["id"]] = ("ok", row)
    if any(v[0] == "ok" for v in out.values()):
        backup(path, min_age=600)
        fd, tmp = tempfile.mkstemp(suffix=os.path.splitext(path)[1], dir=os.path.dirname(path)); os.close(fd)
        try:
            wb.save(tmp); os.replace(tmp, path)
        except PermissionError: raise WorkbookError("The workbook is open in Excel.")
        finally:
            if os.path.exists(tmp):
                try: os.remove(tmp)
                except OSError: pass
    return out

def make_sample(path):
    """A practice workbook laid out like the real payee sheets, so the app can be tested safely."""
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    fills = {"ticket": "FFFFFF", "payref": "FFFFFF", "taxable": "F8D7DA", "cash": "D9EAD3", "check": "FFF2CC", "credit": "D9D9D9"}
    thin = Side(style="thin", color="999999")
    for name, title, rows in (("RVP", "RIO VALLEY PIPE 2026", [("9/30", "OCTOPUS POOLS", 34561, 5664, 585.36), ("9/30", "RODOLFO GUERRA", 34560, 5664, 696.86)]),
                              ("VCP", "VALLEY CONCRETE PRODUCTS 2026", []), ("Magic Valley Concrete", "MAGIC VALLEY CONCRETE 2026", [("10/2", "TOVAR & CRUZ", 34564, 1238, 30666.83)]),
                              ("FORMAT", "TEMPLATE", []), ("DEPOSIT TAXES", "", [])):
        ws = wb.create_sheet(name)
        if name in ("FORMAT", "DEPOSIT TAXES"): continue
        ws["A1"] = title; ws["A1"].font = Font(size=16, bold=True)
        hdr = ["CUSTOMER", "TICKET REF", "PAYMENT REF:", "TAXABLE", "CASH", "CHECK", "CREDIT", "TOTAL", "DEPOSIT DATE", "DEPOSIT AMOUNT"]
        for i, h in enumerate(hdr, 1):
            c = ws.cell(3, i, h); c.font = Font(bold=True); c.alignment = Alignment(horizontal="center", wrap_text=True)
            c.fill = PatternFill("solid", fgColor="C9DAF8")
        for i, w in enumerate([34, 12, 14, 10, 12, 16, 12, 12, 14, 16], 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        for r, (d, cust, tk, pay, amt) in enumerate(rows, 4):
            ws.cell(r, 1, f"{d} {cust}"); ws.cell(r, 2, tk); ws.cell(r, 3, pay)
            ws.cell(r, 6, amt).number_format = "#,##0.00"
            for c, k in enumerate(["customer", "ticket", "payref", "taxable", "cash", "check", "credit"], 1):
                cell = ws.cell(r, c); cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
                if k in fills: cell.fill = PatternFill("solid", fgColor=fills[k])
    wb.save(path)
