"""Admin corrections of payments already sent to Excel: in-place rewrite, moves between day blocks, permissions, safety."""
import os
import openpyxl
from fastapi.testclient import TestClient
from checkflow.excel.layout import Layout

def send(admin, **f):
    base = dict(method="check", customer="EDIT CO", sheet="MVC", amount="100.00", received="2026-11-02", payref="7001")
    p = admin.post("/api/payments", json={"fields": base | f}).json()
    admin.patch(f"/api/payments/{p['id']}", json={"status": "validated"})
    assert admin.post("/api/submit", json={"ids": [p["id"]]}).json()["written"] == 1
    return p["id"]

def ledger_row(admin, pid):
    return next(x for x in admin.get("/api/ledger").json()["rows"] if x["id"] == pid)

def sheet(admin, name):
    ws = openpyxl.load_workbook(admin.get("/api/settings").json()["workbook"])[name]
    return ws, Layout(ws)

def test_same_day_edit_rewrites_the_row_in_place(admin):
    admin.post("/api/settings/practice")
    a = send(admin, customer="EDIT CO", amount="100.00", payref="7001")
    b = send(admin, customer="OTHER CO", amount="50.00", payref="7002")
    r0 = ledger_row(admin, a)["row"]
    res = admin.patch(f"/api/ledger/{a}", json=({"customer": "Edit  Company", "amount": "1,250.5", "payref": "7009", "taxable": "Yes"})).json()
    assert res["excel"] == "synced" and res["row"] == r0 and set(res["changed"]) >= {"customer", "amount", "payref", "taxable"}
    ws, lay = sheet(admin, "MVC")
    assert ws.cell(r0, lay.cols["customer"]).value == "EDIT COMPANY"
    assert ws.cell(r0, lay.amount_cols["check"]).value == 1250.5 and ws.cell(r0, lay.cols["payref"]).value == 7009
    f = ledger_row(admin, a)["fields"]
    assert f["customer"] == "EDIT COMPANY" and f["amount"] == "1250.50" and f["check_no"] == "7009"
    assert ledger_row(admin, b)["row"] != r0                                          # the other payment is untouched

def test_method_change_moves_within_the_day_and_keeps_others_located(admin):
    admin.post("/api/settings/practice")
    a = send(admin, customer="MOVER CO", amount="300.00", payref="8001")
    b = send(admin, customer="STAYER CO", amount="40.00", payref="8002")
    res = admin.patch(f"/api/ledger/{a}", json=({"method": "cash", "payref": "", "ticket": "35999"})).json()
    assert res["excel"] == "synced"
    ws, lay = sheet(admin, "MVC")
    ra, rb = ledger_row(admin, a)["row"], ledger_row(admin, b)["row"]
    assert ws.cell(ra, lay.cols["customer"]).value == "MOVER CO" and ws.cell(ra, lay.amount_cols["cash"]).value == 300
    assert ws.cell(rb, lay.cols["customer"]).value == "STAYER CO"                     # re-found after the block was rebuilt

def test_date_change_empties_the_old_row_and_writes_the_new_day(admin):
    admin.post("/api/settings/practice")
    a = send(admin, customer="LATE CO", amount="75.00", payref="9001", received="2026-11-03")
    old = ledger_row(admin, a)["row"]
    res = admin.patch(f"/api/ledger/{a}", json=({"received": "2026-11-04", "sheet": "RVP"})).json()
    assert res["excel"] == "synced" and res["sheet"] == "RVP"
    mvc, lay = sheet(admin, "MVC")
    assert mvc.cell(old, lay.cols["customer"]).value is None and mvc.cell(old, lay.amount_cols["check"]).value is None
    rio, rlay = sheet(admin, "RVP")
    r = ledger_row(admin, a)
    assert r["sheet"] == "RVP" and rio.cell(r["row"], rlay.cols["customer"]).value == "LATE CO"

def test_only_admins_and_only_valid_values(admin, client):
    admin.post("/api/settings/practice")
    a = send(admin, customer="RULES CO", payref="9101")
    admin.post("/api/users", json={"username": "ar_edit", "password": "arpass1", "role": "ar"})
    ar = TestClient(client.app); ar.post("/api/login", json={"username": "ar_edit", "password": "arpass1"})
    assert ar.patch(f"/api/ledger/{a}", json={"amount": "1"}).status_code == 403
    assert ar.patch(f"/api/ledger/{a}", json={"taxable": "No"}).json()["excel"] == "synced"
    assert admin.patch(f"/api/ledger/{a}", json=({"amount": "abc"})).status_code == 400
    assert admin.patch(f"/api/ledger/{a}", json=({"taxable": "Maybe"})).status_code == 400
    assert admin.patch(f"/api/ledger/{a}", json=({"sheet": "NOPE"})).status_code == 400
    assert admin.patch(f"/api/ledger/{a}", json={"customer": "RULES CO"}).json()["changed"] == []
    assert admin.patch(f"/api/ledger/{a}", json={"routing": "1"}).status_code == 400

def test_open_workbook_queues_the_edit_until_it_is_closed(admin, app_module):
    admin.post("/api/settings/practice")
    a = send(admin, customer="LOCKED CO", payref="9201")
    r0 = ledger_row(admin, a)["row"]
    path = admin.get("/api/settings").json()["workbook"]
    lock = os.path.join(os.path.dirname(path), "~$" + os.path.basename(path))
    open(lock, "w").close()
    try:
        res = admin.patch(f"/api/ledger/{a}", json={"amount": "999.00", "customer": "LOCKED COMPANY"}).json()
        assert res["excel"] == "pending" and "open in Excel" in res["message"]
        assert ledger_row(admin, a)["fields"]["amount"] == "999.00" and ledger_row(admin, a)["pending"]
        res = admin.patch(f"/api/ledger/{a}", json={"invoice": "INV-1"}).json()          # non-Excel field: still queued
        assert res["excel"] == "pending" and ledger_row(admin, a)["pending"]
    finally:
        os.remove(lock)
    assert app_module.push_to_excel() == (0, "")                 # what the 30 s loop does once the workbook is closed
    ws, lay = sheet(admin, "MVC")
    assert not ledger_row(admin, a)["pending"] and ledger_row(admin, a)["row"] == r0
    assert ws.cell(r0, lay.cols["customer"]).value == "LOCKED COMPANY" and ws.cell(r0, lay.amount_cols["check"]).value == 999

def test_imported_rows_are_edited_in_checkflow_only(admin):
    import io, datetime as dt
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Date", "Customer", "Amount", "Payee"]); ws.append([dt.date(2022, 5, 1), "OLD IMPORT CO", 10, "MVC"])
    buf = io.BytesIO(); wb.save(buf)
    pv = admin.post("/api/import/excel", files={"file": ("old.xlsx", buf.getvalue())}).json()
    admin.post(f"/api/import/excel/{pv['token']}", json={})
    pid = next(x["id"] for x in admin.get("/api/ledger").json()["rows"] if x["fields"]["customer"] == "OLD IMPORT CO")
    res = admin.patch(f"/api/ledger/{pid}", json=({"sheet": "RVP", "amount": "12.00"})).json()
    assert res["excel"] == "skipped" and ledger_row(admin, pid)["sheet"] == "RVP" and ledger_row(admin, pid)["fields"]["amount"] == "12.00"
