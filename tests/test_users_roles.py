"""User management and the three roles: admin, ar (read + download images + Taxable), readonly."""
import io
from fastapi.testclient import TestClient
from PIL import Image

def login(client, who, pw):
    c = TestClient(client.app)
    r = c.post("/api/login", json={"username": who, "password": pw})
    return c, r

def sent_payment(admin, **f):
    base = dict(method="check", customer="ROLE TEST CO", sheet="MVC", amount="200.00", received="2026-11-10", payref="5501")
    p = admin.post("/api/payments", json={"fields": base | f}).json()
    admin.patch(f"/api/payments/{p['id']}", json={"status": "validated"})
    assert admin.post("/api/submit", json={"ids": [p["id"]]}).json()["written"] == 1
    return p["id"]

def jpg():
    b = io.BytesIO(); Image.new("RGB", (300, 400), "white").save(b, "JPEG"); return b.getvalue()

def test_bulk_add_and_first_sign_in(admin, client):
    text = ("Sandy Test, sandy.t@example.com, Admin\n"
            "Michelle Test\tmichelle.t@example.com\tAR: read, download images, taxes\n"
            "Richard Test   [richard.t@example.com](mailto:richard.t@example.com)   Read only\n"
            "Dup Person, sandy.t@example.com, Admin\n")
    res = admin.post("/api/users/bulk", json={"text": text}).json()["results"]
    assert [(r["email"], r["role"], r["ok"]) for r in res] == [("sandy.t@example.com", "admin", True), ("michelle.t@example.com", "ar", True),
                                                               ("richard.t@example.com", "readonly", True), ("sandy.t@example.com", "admin", False)]
    assert res[2]["name"] == "Richard Test" and res[0]["password"]
    # first sign-in with the email (any case) and the temporary password -> must choose a new one
    c, r = login(client, "Michelle.T@Example.com", res[1]["password"])
    assert r.status_code == 200
    me = c.get("/api/me").json()["user"]
    assert me["must_change"] and me["role"] == "ar" and me["email"] == "michelle.t@example.com"
    assert c.post("/api/me/password", json={"current": res[1]["password"], "new": res[1]["password"]}).status_code == 400
    assert c.post("/api/me/password", json={"current": res[1]["password"], "new": "abc"}).status_code == 200   # any password they like
    assert not c.get("/api/me").json()["user"]["must_change"]
    assert login(client, "michelle.t@example.com", "abc")[1].status_code == 200

def test_forgot_password_request_and_admin_reset(admin, client):
    r = admin.post("/api/users", json={"email": "lost.pw@example.com", "display_name": "Lost Pw", "role": "readonly"}).json()
    uid = next(u["id"] for u in r["users"] if u["email"] == "lost.pw@example.com")
    anon = TestClient(client.app)
    assert anon.post("/api/password-reset-request", json={"login": "lost.pw@example.com"}).json()["ok"]
    assert anon.post("/api/password-reset-request", json={"login": "nobody@example.com"}).json()["ok"]   # same answer
    reqs = admin.get("/api/password-requests").json()
    assert [q["email"] for q in reqs] == ["lost.pw@example.com"]
    out = admin.post(f"/api/users/{uid}/password", json={}).json()                                     # generated
    assert out["password"] and admin.get("/api/password-requests").json() == []
    c, r = login(client, "lost.pw@example.com", out["password"])
    assert r.status_code == 200 and c.get("/api/me").json()["user"]["must_change"]
    out = admin.post(f"/api/users/{uid}/password", json={"password": "typed-by-admin", "must_change": False}).json()
    c, r = login(client, "lost.pw@example.com", "typed-by-admin")
    assert r.status_code == 200 and not c.get("/api/me").json()["user"]["must_change"]

def test_roles_are_enforced_on_the_server(admin, client):
    admin.post("/api/settings/practice")
    pid = sent_payment(admin)
    page = admin.post(f"/api/payments/{pid}/attachments", files={"files": ("receipt.jpg", jpg(), "image/jpeg")}).json()["pages"][0]["id"]
    users = {}
    for role in ("ar", "readonly"):
        pw = admin.post("/api/users", json={"email": f"{role}.role@example.com", "role": role, "password": "pw1"}).json()["password"]
        users[role] = login(client, f"{role}.role@example.com", pw)[0]
    ar, ro = users["ar"], users["readonly"]
    # both read the ledger and view images
    for c in (ar, ro):
        assert c.get("/api/ledger").status_code == 200 and c.get(f"/api/pages/{page}").status_code == 200
        assert c.post("/api/ledger/export", json={"ids": [pid]}).status_code == 200
    # downloads: AR yes, read-only no
    d = ar.get(f"/api/pages/{page}?download=1")
    assert d.status_code == 200 and "ROLE-TEST-CO" in d.headers["content-disposition"]
    assert ar.get(f"/api/payments/{pid}/document.pdf").headers["content-type"] == "application/pdf"
    assert ro.get(f"/api/pages/{page}?download=1").status_code == 403
    assert ro.get(f"/api/payments/{pid}/document.pdf").status_code == 403
    # Taxable: AR yes, read-only no; other fields and admin pages: neither
    assert ar.patch(f"/api/ledger/{pid}", json={"taxable": "Yes"}).status_code == 200
    assert ro.patch(f"/api/ledger/{pid}", json={"taxable": "No"}).status_code == 403
    # attachments: AR can add and remove them (never scanned pages), read-only can't
    added = ar.post(f"/api/payments/{pid}/attachments", files={"files": ("auth.jpg", jpg(), "image/jpeg")})
    assert added.status_code == 200 and len(added.json()["pages"]) == 2
    assert ar.delete(f"/api/payments/{pid}/attachments/{added.json()['pages'][-1]['id']}").status_code == 200
    assert ro.post(f"/api/payments/{pid}/attachments", files={"files": ("x.jpg", jpg(), "image/jpeg")}).status_code == 403
    assert ro.delete(f"/api/payments/{pid}/attachments/{page}").status_code == 403
    for c in (ar, ro):
        assert c.patch(f"/api/ledger/{pid}", json={"amount": "1"}).status_code == 403
        assert c.post(f"/api/ledger/{pid}/delete", json={}).status_code == 403
        assert c.get("/api/users").status_code == 403 and c.get("/api/analytics").status_code == 403

def test_last_admin_is_protected(admin):
    me = next(u for u in admin.get("/api/users").json() if u["username"] == "admin")
    assert admin.patch(f"/api/users/{me['id']}", json={"role": "readonly"}).status_code == 400
    assert admin.patch(f"/api/users/{me['id']}", json={"active": False}).status_code == 400

def test_posted_to_netsuite_and_who_did_it(admin, client):
    admin.post("/api/settings/practice")
    pid = sent_payment(admin, customer="POSTED CO")
    admin.post("/api/users", json={"email": "post.ar@example.com", "display_name": "Michelle Test", "role": "ar", "password": "pw"})
    admin.post("/api/users", json={"email": "post.ro@example.com", "role": "readonly", "password": "pw"})
    ar = login(client, "post.ar@example.com", "pw")[0]; ro = login(client, "post.ro@example.com", "pw")[0]
    assert ar.patch(f"/api/ledger/{pid}", json={"taxable": "Yes"}).status_code == 200
    r = ar.patch(f"/api/ledger/{pid}", json={"posted": " NS-48213 "}).json()
    assert r["excel"] == "unchanged"                                                  # NetSuite status isn't written to Excel
    f = next(x for x in admin.get("/api/ledger").json()["rows"] if x["id"] == pid)["fields"]
    assert f["posted"] == "NS-48213" and f["posted_by"] == "Michelle Test" and f["posted_at"][:2] == "20"
    assert f["taxable_by"] == "Michelle Test" and f["taxable_at"]
    assert ro.patch(f"/api/ledger/{pid}", json={"posted": ""}).status_code == 403
    assert ar.patch(f"/api/ledger/{pid}", json={"posted": "X" * 41}).status_code == 400
    assert ar.patch(f"/api/ledger/{pid}", json={"amount": "1", "posted": ""}).status_code == 403
    # admin un-posts: the status and the person change, and the history keeps both
    admin.patch(f"/api/ledger/{pid}", json={"posted": ""})
    f = next(x for x in admin.get("/api/ledger").json()["rows"] if x["id"] == pid)["fields"]
    assert f["posted"] == "" and f["posted_by"] == "admin"
    h = ro.get(f"/api/ledger/{pid}/history").json()
    edits = [(x["name"], x["detail"]) for x in h if x["action"] == "ledger_edit"]
    assert ("admin", {"posted": ["NS-48213", ""]}) in edits and ("Michelle Test", {"posted": ["", "NS-48213"]}) in edits
    assert ("Michelle Test", {"taxable": ["", "Yes"]}) in edits and any(x["action"] == "submitted" or x["action"] for x in h)
    # exports carry who / when
    import openpyxl
    ar.patch(f"/api/ledger/{pid}", json={"posted": "NS-50001"})
    ws = openpyxl.load_workbook(io.BytesIO(admin.post("/api/ledger/export", json={"ids": [pid]}).content))["Ledger"]
    head = [c.value for c in ws[1]]; row = {h: c.value for h, c in zip(head, ws[2])}
    assert row["NetSuite #"] == "NS-50001" and row["Posted by"] == "Michelle Test" and row["Taxable set by"] == "Michelle Test"
