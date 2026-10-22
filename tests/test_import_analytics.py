"""Admin features: spreadsheet import into the ledger (preview -> commit -> undo) and analytics (KPIs, projection, storylines)."""
import datetime as dt, io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from checkflow import analytics
from checkflow.excel import ledger_import as li
from checkflow.registry import StaticRegistry

# ------------------------------------------------------------------ reading files
def rec(d, cust="ACME", amt="100.00", src="S"):
    return dict(sheet="MVC", row=1, received=d, customer=cust, amount=amt, source=src)

def test_flat_table_headers_methods_and_payees():
    rows = [["Monthly report"], [], ["Date Received", "Client Name", "Payment Type", "Check #", "Amount", "Payee", "Taxable"],
            [dt.datetime(2026, 1, 5), "Acme  Builders", "Visa", "", 1250.5, "Magic Valley Concrete", "Y"],
            ["01/06/2026", "Juan Cantu", "ACH", "778", "$2,000.00", "RVP", "no"],
            ["2026-01-07", "Bo Pools", "", "1001", "300", "", ""],
            ["2026-01-08", "", "Cash", "", "50", "MVC", ""]]                       # no customer -> skipped
    reg = StaticRegistry()
    resolve = lambda t: (reg.match_payee(t) or (None, ""))[1] or ({"RVP": "RVP", "MVC": "MVC"}.get(t.upper(), ""))
    recs, skipped = li.table_records(rows, "Sheet1", resolve)
    assert skipped == 1 and [r["method"] for r in recs] == ["card", "ach", "check"]
    assert recs[0]["customer"] == "ACME BUILDERS" and recs[0]["sheet"] == "MVC" and recs[0]["taxable"] == "Yes" and recs[0]["received"] == "2026-01-05"
    assert recs[1]["sheet"] == "RVP" and recs[1]["amount"] == "2000.00" and recs[1]["taxable"] == "No" and recs[1]["payref"] == "778"
    assert recs[2]["sheet"] == ""                                                    # asks for a payee in the preview

def test_csv_file_keeps_negative_rows(tmp_path):
    p = tmp_path / "pay.csv"
    p.write_text("date,customer,amount,method,payee\n2026-02-01,ACME,100,check,MVC\n2026-02-01,WITHDRAWAL,-50,,MVC\n"
                 "2026-02-01,CASH ON HAND,(25.50),cash,MVC\n2026-02-02,BETA,0,,MVC\n2026-02-03,WITHDRAWAL,40,,MVC\n")
    res = li.read_file(str(p), "pay.csv", ["MVC"], StaticRegistry())
    assert [(r["customer"], r["amount"]) for r in res["records"]] == [("ACME", "100.00"), ("WITHDRAWAL", "-50.00"),
                                                                       ("CASH ON HAND", "-25.50"), ("WITHDRAWAL", "40.00")]
    assert any("2 negative rows" in w and "-$75.50" in w for w in res["warnings"])
    assert any("positive amount" in w and "check its sign" in w for w in res["warnings"])
    assert res["skipped"] == 1                                                       # the zero row

def test_year_typo_is_flagged_and_fixed():
    ds = ["2025-07-14", "2025-07-15", "2025-07-15", "2026-07-16", "2026-07-16", "2025-07-17", "2025-07-18", "2025-10-31", "2025-07-21", "2025-07-22"]
    recs = li.flag_dates([rec(d) for d in ds])
    flagged = {i: r["date_issue"] for i, r in enumerate(recs) if r.get("date_issue")}
    assert flagged[3]["fix"] == "2025-07-16" and flagged[4]["fix"] == "2025-07-16"
    assert flagged[7]["fix"] is None                                                 # not a year typo: kept for a person to check
    assert 5 not in flagged and 6 not in flagged                                     # neighbours of the typos are fine
    out, fixed, dropped = li.apply_date_policy(recs, "fix")
    assert fixed == 2 and dropped == 0 and len(out) == len(ds) and out[3]["received"] == "2025-07-16" and out[7]["received"] == "2025-10-31"
    assert li.apply_date_policy(recs, "skip")[2] == 3 and li.apply_date_policy(recs, "keep")[0][3]["received"] == "2026-07-16"

def test_duplicates_count_like_a_multiset():
    recs = [rec("2026-01-05"), rec("2026-01-05"), rec("2026-01-06")]
    li.mark_duplicates(recs, {li.fingerprint("MVC", "2026-01-05", "100.00", "acme"): 1})
    assert [r["duplicate"] for r in recs] == [True, False, False]                    # two same-day payments stay two

# ------------------------------------------------------------------ API: preview, commit, re-upload, undo
def table_xlsx(rows):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Export"
    ws.append(["Date", "Customer", "Amount", "Method", "Payee"])
    for r in rows: ws.append(r)
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

def test_import_preview_commit_dedupe_undo(admin, client):
    data = table_xlsx([[dt.date(2024, 3, 1), "IMPORT CO", 1000, "check", "TST"], [dt.date(2024, 3, 2), "IMPORT CO", 500, "wire", "TST"],
                       [dt.date(2024, 3, 3), "OTHER LLC", 250, "card", ""]])
    pv = admin.post("/api/import/excel", files={"file": ("export.xlsx", data)}).json()
    assert pv["rows"] == 3 and pv["new"] == 3 and pv["needs_sheet"] and pv["formats"] == {"Export": "table"}
    assert admin.post(f"/api/import/excel/{pv['token']}", json={}).status_code == 400          # rows without a payee need one
    r = admin.post(f"/api/import/excel/{pv['token']}", json={"default_sheet": "tst"}).json()
    assert r["imported"] == 3
    pv2 = admin.post("/api/import/excel", files={"file": ("again.xlsx", data)}).json()
    assert pv2["new"] == 0 and pv2["duplicates"] == 3
    rows = [x for x in admin.get("/api/ledger").json()["rows"] if x["fields"]["customer"] in ("IMPORT CO", "OTHER LLC")]
    assert len(rows) == 3 and all(x["imported"] and x["sheet"] == "TST" and x["row"] is None for x in rows)
    # a ledger edit on an uploaded row never queues a workbook write
    res = admin.patch(f"/api/ledger/{rows[0]['id']}", json={"taxable": "Yes"}).json()
    assert res["excel"] == "skipped"
    from checkflow import db
    assert not any(d["id"] == rows[0]["id"] for d in db.dirty_rows())
    # AR can't import; undo removes the batch only
    admin.post("/api/users", json={"username": "ar_imp", "password": "arpass1", "role": "ar"})
    ar = TestClient(client.app); ar.post("/api/login", json={"username": "ar_imp", "password": "arpass1"})
    assert ar.post("/api/import/excel", files={"file": ("x.xlsx", data)}).status_code == 403
    assert ar.get("/api/analytics").status_code == 403
    b = next(x for x in admin.get("/api/import/batches").json() if x["filename"] == "export.xlsx")
    assert b["rows"] == 3 and b["total"] == 1750
    assert admin.delete(f"/api/import/batches/{b['batch']}").json()["removed"] == 3
    assert not [x for x in admin.get("/api/ledger").json()["rows"] if x["fields"]["customer"] == "IMPORT CO"]

def test_rejects_other_file_types(admin):
    assert admin.post("/api/import/excel", files={"file": ("old.xls", b"x")}).status_code == 400
    assert admin.post("/api/import/excel", files={"file": ("junk.xlsx", b"not a zip")}).status_code == 400

# ------------------------------------------------------------------ analytics
def test_periods_and_business_days():
    d = dt.date(2026, 10, 14)
    assert analytics.pkey(d, "week") == dt.date(2026, 10, 12) and analytics.pkey(d, "quarter") == dt.date(2026, 10, 1)
    assert analytics.pnext(dt.date(2026, 11, 1), "quarter") == dt.date(2027, 2, 1)
    assert analytics.bdays(dt.date(2026, 10, 12), dt.date(2026, 10, 18)) == 5

def test_forecast_follows_a_straight_trend():
    hist = [(dt.date(2025, m, 1), 1000.0 + 100 * m) for m in range(1, 13)]
    fc = analytics.forecast(hist, "month", 3, dt.date(2025, 12, 1))
    assert [p["key"] for p in fc["points"]] == ["2026-01-01", "2026-02-01", "2026-03-01"]
    assert fc["points"][0]["value"] == pytest.approx(2300, abs=1) and fc["points"][0]["low"] <= 2300 <= fc["points"][0]["high"]
    assert analytics.forecast(hist[:3], "month", 3, dt.date(2025, 3, 1)) is None              # too little history

def test_analytics_endpoint(admin):
    rows = []
    for m in range(1, 13):                                    # ANA sheet: two clients, steady growth through 2023
        rows.append([dt.date(2023, m, 10), "STEADY CO", 1000 * m, "check", "ANA"])
        if m <= 6: rows.append([dt.date(2023, m, 20), "QUIET CO", 400, "card", "ANA"])
    pv = admin.post("/api/import/excel", files={"file": ("ana.xlsx", table_xlsx(rows))}).json()
    admin.post(f"/api/import/excel/{pv['token']}", json={})
    a = admin.get("/api/analytics", params={"start": "2023-07-01", "end": "2023-12-31", "grain": "month", "sheets": "ANA"}).json()
    k = a["kpis"]
    assert k["current"]["total"] == sum(1000 * m for m in range(7, 13)) and k["current"]["count"] == 6
    assert k["previous"]["total"] == sum(1000 * m for m in range(1, 7)) + 6 * 400
    assert [p["label"] for p in a["periods"]] == ["Jul 2023", "Aug 2023", "Sep 2023", "Oct 2023", "Nov 2023", "Dec 2023"]
    assert a["series"][0]["key"] == "ANA" and len(a["projection"]["points"]) == 3
    assert a["projection"]["points"][0]["label"] == "Jan 2024"
    ids = {s["id"] for s in a["storylines"]}
    assert {"headline", "forecast", "lapsed"} <= ids                       # QUIET CO paid monthly, then stopped
    assert "QUIET CO".title() in next(s for s in a["storylines"] if s["id"] == "lapsed")["text"]
    one = admin.get("/api/analytics", params={"customer": "steady co", "start": "2023-01-01", "end": "2023-12-31", "sheets": "ANA"}).json()
    assert one["kpis"]["current"]["count"] == 12 and any(s["id"] == "client-cadence" for s in one["storylines"])
    bd = admin.get("/api/analytics", params={"by": "customer", "start": "2023-01-01", "end": "2023-12-31", "sheets": "ANA"}).json()["breakdown"]
    assert [r["key"] for r in bd] == ["STEADY CO", "QUIET CO"] and bd[0]["share"] > 90

def test_negative_rows_import_into_ledger_not_analytics_and_full_export(admin):
    data = table_xlsx([[dt.date(2023, 5, 1), "NEG CO", 900, "check", "NEG"], [dt.date(2023, 5, 1), "WITHDRAWAL", -300, "cash", "NEG"],
                       [dt.date(2023, 5, 2), "NEG CO", -120.5, "card", "NEG"]])
    pv = admin.post("/api/import/excel", files={"file": ("neg.xlsx", data)}).json()
    assert pv["rows"] == 3 and any("2 negative rows" in w for w in pv["warnings"])
    assert admin.post(f"/api/import/excel/{pv['token']}", json={}).json()["imported"] == 3
    led = [x for x in admin.get("/api/ledger").json()["rows"] if x["sheet"] == "NEG"]
    assert sorted(float(x["fields"]["amount"]) for x in led) == [-300, -120.5, 900]
    # analytics: money received only
    an = admin.get("/api/analytics", params={"start": "2023-05-01", "end": "2023-05-02", "by": "customer"}).json()
    assert an["kpis"]["current"]["total"] == 900 and an["kpis"]["current"]["count"] == 1
    # Export Excel with no ids = the whole ledger (negatives included); with ids = just those
    whole = openpyxl.load_workbook(io.BytesIO(admin.post("/api/ledger/export", json={}).content))["Ledger"]
    n_all = len(admin.get("/api/ledger").json()["rows"])
    assert whole.max_row == n_all + 2                                                # header + rows + total
    assert -300 in [whole.cell(r, 8).value for r in range(2, whole.max_row)]
    some = openpyxl.load_workbook(io.BytesIO(admin.post("/api/ledger/export", json={"ids": [led[0]["id"]]}).content))["Ledger"]
    assert some.max_row == 3
