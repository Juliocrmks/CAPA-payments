"""Excel writer: day blocks in the sheet's own layout, same-day merge, formulas, edits, locks, import."""
import os, datetime
import openpyxl
import pytest
from checkflow.excel import writer, importer
from checkflow.excel.layout import Layout
from conftest import REFS

def item(i, method, amount, customer="ACME BUILDERS", sheet="MVC", day="2026-10-09", **kw):
    return dict(id=i, sheet=sheet, received=day, customer=customer, order=kw.get("order", ""), ticket=kw.get("ticket", ""),
                payref=kw.get("payref", ""), method=method, amount=amount, taxable=kw.get("taxable", ""), deposit_bank=kw.get("bank", ""))

def block(ws, start):
    out = []
    for r in range(start, ws.max_row + 1):
        vals = {c: ws.cell(r, c).value for c in range(1, 15) if ws.cell(r, c).value not in (None, "")}
        if vals: out.append((r, vals))
    return out

def test_layout_detects_practice_columns(workbook):
    lay = Layout(openpyxl.load_workbook(workbook)["UPPER"])
    assert lay.cols["customer"] == 2 and lay.cols["order"] == 3 and lay.cols["taxable"] == 6 and lay.cols["credit"] == 9
    mvc = Layout(openpyxl.load_workbook(workbook)["MVC"])
    assert mvc.labels == {"check": "CHECK", "cash": "CASH", "credit": "CREDIT"} and mvc.order[:2] == ["credit", "check"]

def test_write_new_day_block(workbook):
    lay0 = Layout(openpyxl.load_workbook(workbook)["MVC"])
    rows = writer.write_payments(workbook, [item(1, "check", "100.50", payref="1234", bank="TNB", taxable="Yes"),
                                            item(2, "card", "20.00", payref="6759", ticket="35391"),
                                            item(3, "wire", "300.00", customer="ESPERANZA HOMES")])
    ws = openpyxl.load_workbook(workbook)["MVC"]
    lay = Layout(ws)
    assert lay.hdr > lay0.hdr                               # a new header row for the new day
    r1, r2, r3 = rows[1][1], rows[2][1], rows[3][1]
    assert ws.cell(r1, 1).value == "OCT/09/26" and ws.cell(r1, lay.cols["check"]).value == 100.5
    assert ws.cell(r1, lay.cols["taxable"]).value == "YES" and ws.cell(r1, lay.cols["bank"]).value == "TNB"
    assert ws.cell(r2, lay.cols["credit"]).value == 20.0 and ws.cell(r2, lay.cols["ticket"]).value == 35391
    assert ws.cell(r3, lay.cols["ticket"]).value == "WIRE" and isinstance(ws.cell(r3, lay.cols["payref"]).value, datetime.datetime)
    assert r2 < r1                                         # this sheet lists CREDIT before CHECK
    formulas = [(r, v) for r, vals in block(ws, lay.hdr) for c, v in vals.items() if isinstance(v, str) and v.startswith("=")]
    labels = [ws.cell(r, lay.label_col).value for r, _ in formulas]
    assert labels == ["CREDIT", "CHECK", "TOTAL"]
    assert formulas[0][1] == f"=SUM(H{r2}:H{formulas[0][0] - 1})"

def test_same_day_second_batch_merges_and_keeps_rows(workbook):
    writer.write_payments(workbook, [item(1, "check", "100.00", payref="1"), item(2, "card", "5.00", payref="2")])
    ws = openpyxl.load_workbook(workbook)["MVC"]; ws.cell(Layout(ws).hdr + 1, 2).value  # sanity
    rows = writer.write_payments(workbook, [item(3, "check", "200.00", payref="3", customer="SECOND BATCH")])
    ws = openpyxl.load_workbook(workbook)["MVC"]
    lay = Layout(ws)
    customers = [ws.cell(r, 2).value for r in range(lay.hdr + 1, ws.max_row + 1) if ws.cell(r, 2).value]
    assert customers.count("ACME BUILDERS") == 2 and "SECOND BATCH" in customers
    assert sum(1 for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "OCT/09/26") == 3   # still one block for the day
    sub = next(r for r in range(lay.hdr, ws.max_row + 1) if ws.cell(r, lay.label_col).value == "CHECK")
    first = next(r for r in range(lay.hdr, ws.max_row + 1) if ws.cell(r, lay.cols["check"]).value == 100.0)
    assert ws.cell(sub, lay.sum_col).value == f"=SUM(G{first}:G{sub - 1})" and rows[3][1] < sub

def test_other_days_untouched(workbook):
    before = [[c.value for c in row] for row in openpyxl.load_workbook(workbook)["RIO"].iter_rows()]
    writer.write_payments(workbook, [item(1, "check", "1.00", sheet="RIO")])
    after = [[c.value for c in row] for row in openpyxl.load_workbook(workbook)["RIO"].iter_rows()]
    assert after[:len(before)] == before

def test_lock_file_blocks_writes(workbook):
    lock = writer.lock_path(workbook)
    open(lock, "w").close()
    try:
        with pytest.raises(writer.WorkbookError, match="open in Excel"):
            writer.write_payments(workbook, [item(1, "check", "1.00")])
    finally:
        os.remove(lock)

def test_backup_written(workbook):
    writer.write_payments(workbook, [item(1, "check", "1.00")])
    assert os.listdir(os.path.join(os.path.dirname(workbook), "checkflow_backups"))

def test_update_rows_finds_moved_row(workbook):
    rows = writer.write_payments(workbook, [item(1, "check", "10.00", customer="MOVER", payref="9")])
    res = writer.update_rows(workbook, [dict(id=1, sheet="MVC", row=rows[1][1] + 5, received="2026-10-09", customer="MOVER",
                                             amount="10.00", taxable="No", deposit_bank="LSB")])
    assert res[1] == ("ok", rows[1][1])
    ws = openpyxl.load_workbook(workbook)["MVC"]; lay = Layout(ws)
    assert ws.cell(rows[1][1], lay.cols["taxable"]).value == "NO" and ws.cell(rows[1][1], lay.cols["bank"]).value == "LSB"

def test_unknown_sheet(workbook):
    with pytest.raises(writer.WorkbookError, match="does not exist"):
        writer.write_payments(workbook, [item(1, "check", "1.00", sheet="NOPE")])

def test_importer_reads_back(workbook):
    writer.write_payments(workbook, [item(1, "wire", "77.00", customer="WIRED INC")])
    recs = list(importer.history(workbook, ["MVC"]))
    w = next(r for r in recs if r["customer"] == "WIRED INC")
    assert w["method"] == "wire" and w["amount"] == "77.00" and w["received"] == "2026-10-09"
    assert "TOVAR & CRUZ" in importer.customers(workbook)

REAL = os.path.join(REFS, "DAILY TRANSACTION 2025 NEW.xlsx")

@pytest.mark.skipif(not os.path.exists(REAL), reason="real workbook not in references/")
def test_real_workbook_layouts_and_totals(tmp_path):
    """The 2025 workbook: every payee sheet is understood, and a written day sums like the printed 10/09 sheets."""
    import json, shutil
    p = str(tmp_path / "real.xlsx"); shutil.copy(REAL, p)
    wb = openpyxl.load_workbook(p)
    for s in ("MVC", "UPPER", "RIO", "VCP"):
        lay = Layout(wb[s]); assert {"customer", "check", "cash", "credit", "payref"} <= set(lay.cols)
    key = json.load(open(os.path.join(os.path.dirname(__file__), "fixtures", "answer_key_2026-10-09.json")))
    items = [item(n, x["method"], x["amount"], customer=x["customer"], payref=x["payref"], ticket=x["ticket"])
             for n, x in enumerate(key["MVC OCT_09-26.pdf"]["payments"], 1)]
    rows = writer.write_payments(p, items)
    ws = openpyxl.load_workbook(p)["MVC"]; lay = Layout(ws)
    sums = {}
    for r in range(lay.hdr, ws.max_row + 1):
        f = ws.cell(r, lay.sum_col).value
        if isinstance(f, str) and f.startswith("=SUM("):
            col, a, b = f[5], int(f[6:f.index(":")]), int(f[f.index(":") + 2:-1])
            sums[ws.cell(r, lay.label_col).value] = round(sum(ws.cell(x, openpyxl.utils.column_index_from_string(col)).value or 0 for x in range(a, b + 1)), 2)
    assert sums == {"CHECKS": 292790.93, "CREDIT": 56521.58}     # matches the printed MVC sheet for OCT/09/26
    assert len(rows) == 8
