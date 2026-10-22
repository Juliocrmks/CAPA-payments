"""Workbooks: list what's there, create a new (empty) one from the current layout, upload, download, select.
Payments stay tied to the workbook they were written to."""
import io, os, uuid
import openpyxl
import pytest
from checkflow.excel.layout import Layout

def sent(admin, **f):
    base = dict(method="check", customer="BOOK CO", sheet="MVC", amount="410.00", received="2026-11-20", payref="7101")
    p = admin.post("/api/payments", json={"fields": base | f}).json()
    admin.patch(f"/api/payments/{p['id']}", json={"status": "validated"})
    assert admin.post("/api/submit", json={"ids": [p["id"]]}).json()["written"] == 1
    return p["id"]

def name(): return f"Payments {uuid.uuid4().hex[:6]}.xlsx"

@pytest.fixture(autouse=True)
def back_to_practice(admin):
    admin.post("/api/settings/practice")
    yield
    admin.post("/api/settings/practice")

def test_new_workbook_from_current_layout(admin):
    old = admin.get("/api/settings").json()["workbook"]
    n = name()
    r = admin.post("/api/workbooks/new", json={"name": n, "year": 2027}).json()
    assert r["sheets"] == ["MVC", "UPPER", "RIO", "VCP"] and "Sheet1" in r["other"]
    assert r["settings"]["workbook"] == r["path"] and os.path.basename(r["path"]) == n
    books = {b["name"]: b for b in r["workbooks"]}
    assert books[n]["current"] and books["Practice_Workbook.xlsx"]["practice"]
    wb = openpyxl.load_workbook(r["path"])
    ws, src = wb["MVC"], Layout(openpyxl.load_workbook(old)["MVC"])
    lay = Layout(ws)
    assert ws["B1"].value == "MAGIC VALLEY CONCRETE 2027"                          # title year updated
    assert not any(lay.is_entry(x) for x in range(1, ws.max_row + 1))              # no payments carried over
    assert (lay.order, lay.labels, lay.cols) == (src.order, src.labels, src.cols)  # same layout
    # payments now go into the new workbook, styled like the old one
    pid = sent(admin, customer="NEW YEAR CO")
    ws = openpyxl.load_workbook(r["path"])["MVC"]
    row = next(x["row"] for x in admin.get("/api/ledger").json()["rows"] if x["id"] == pid)
    assert ws.cell(row, Layout(ws).cols["customer"]).value == "NEW YEAR CO"
    assert ws.cell(row, Layout(ws).cols["check"]).fill.fgColor.rgb.endswith("FFF2CC")
    # the same name twice is refused
    assert admin.post("/api/workbooks/new", json={"name": n}).status_code == 409

def test_standard_layout_and_keep_current(admin):
    cur = admin.get("/api/settings").json()["workbook"]
    r = admin.post("/api/workbooks/new", json={"name": name(), "source": "standard", "use": False}).json()
    assert r["settings"]["workbook"] == cur and r["sheets"] == ["MVC", "UPPER", "RIO", "VCP"]

def test_edits_go_back_to_the_workbook_the_payment_was_written_to(admin):
    first = admin.get("/api/settings").json()["workbook"]
    pid = sent(admin, customer="LAST YEAR CO", amount="500.00")
    nxt = admin.post("/api/workbooks/new", json={"name": name()}).json()["path"]
    r = admin.patch(f"/api/ledger/{pid}", json={"amount": "550.00"}).json()
    assert r["excel"] == "synced"
    ws = openpyxl.load_workbook(first)["MVC"]; lay = Layout(ws)
    row = next(x for x in admin.get("/api/ledger").json()["rows"] if x["id"] == pid)
    assert row["workbook"] == "Practice_Workbook.xlsx"
    assert ws.cell(row["row"], lay.cols["check"]).value == 550
    new_ws = openpyxl.load_workbook(nxt)["MVC"]
    assert "LAST YEAR CO" not in [new_ws.cell(x, 2).value for x in range(1, new_ws.max_row + 1)]

def test_upload_select_download(admin, workbook):
    data = open(workbook, "rb").read()
    up = admin.post("/api/workbooks/upload", files={"file": ("Daily 2026.xlsx", data, "application/octet-stream")}).json()
    again = admin.post("/api/workbooks/upload", files={"file": ("Daily 2026.xlsx", data, "application/octet-stream")}).json()
    assert os.path.basename(again["path"]).startswith("Daily 2026 (") and up["path"] != again["path"]
    assert admin.post("/api/workbooks/upload", files={"file": ("x.xlsx", b"not excel", "application/octet-stream")}).status_code == 400
    s = admin.post("/api/workbooks/select", json={"path": up["path"]}).json()
    assert s["settings"]["workbook"] == up["path"] and s["settings"]["sheets"][:4] == ["MVC", "UPPER", "RIO", "VCP"]
    d = admin.get("/api/workbooks/download", params={"path": up["path"]})
    assert d.status_code == 200 and d.content == data and "Daily%202026.xlsx" in d.headers["content-disposition"]
    assert admin.get("/api/workbooks/download", params={"path": "/etc/passwd"}).status_code == 404
    assert admin.post("/api/workbooks/select", json={"path": "/etc/hosts"}).status_code == 404

def test_negative_amount_edits(admin):
    pid = sent(admin, customer="REFUND CO", amount="-250.00")
    assert admin.patch(f"/api/ledger/{pid}", json={"amount": "-300"}).status_code == 200
    assert admin.patch(f"/api/ledger/{pid}", json={"amount": "0"}).status_code == 400

def test_workbooks_are_admin_only(admin, client):
    from fastapi.testclient import TestClient
    pw = admin.post("/api/users", json={"email": f"ar.{uuid.uuid4().hex[:5]}@example.com", "role": "ar", "password": "pw"}).json()
    ar = TestClient(client.app)
    email = next(u["email"] for u in pw["users"] if u["role"] == "ar")
    ar.post("/api/login", json={"username": email, "password": "pw"})
    assert ar.get("/api/workbooks").status_code == 403
    assert ar.post("/api/workbooks/new", json={"name": "x"}).status_code == 403
