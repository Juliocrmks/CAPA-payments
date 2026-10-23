"""The ledger is the source of truth: attachments, delete to a Deleted section (out of Excel), restore, Excel export."""
import io
import openpyxl
from PIL import Image
from checkflow.excel.layout import Layout

def sent(admin, **f):
    base = dict(method="check", customer="LIFE CO", sheet="MVC", amount="300.00", received="2026-11-12", payref="6101")
    p = admin.post("/api/payments", json={"fields": base | f}).json()
    admin.patch(f"/api/payments/{p['id']}", json={"status": "validated"})
    assert admin.post("/api/submit", json={"ids": [p["id"]]}).json()["written"] == 1
    return p["id"]

def row(admin, pid):
    return next((x for x in admin.get("/api/ledger").json()["rows"] if x["id"] == pid), None)

def mvc(admin):
    ws = openpyxl.load_workbook(admin.get("/api/settings").json()["workbook"])["MVC"]
    return ws, Layout(ws)

def pdf_bytes(pages=3):
    imgs = [Image.new("RGB", (400, 520), c) for c in ("white", "lightgray", "white")[:pages]]
    b = io.BytesIO(); imgs[0].save(b, "PDF", save_all=True, append_images=imgs[1:]); return b.getvalue()

def test_attachments_added_and_removed(admin):
    admin.post("/api/settings/practice")
    pid = sent(admin)
    p = admin.post(f"/api/payments/{pid}/attachments", files=[("files", ("card.pdf", pdf_bytes(3), "application/pdf")),
                                                              ("files", ("receipt.png", _png(), "image/png"))]).json()
    att = [x for x in p["pages"] if x["kind"] == "attachment"]
    assert len(att) == 4
    pdf = admin.get(f"/api/payments/{pid}/document.pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF" and pdf.content.count(b"/Type /Page") >= 4
    assert len(admin.delete(f"/api/payments/{pid}/attachments/{att[0]['id']}").json()["pages"]) == 3
    assert admin.post(f"/api/payments/{pid}/attachments", files={"files": ("photo.heic", b"x", "image/heic")}).status_code == 400

def _png():
    b = io.BytesIO(); Image.new("RGB", (200, 260), "white").save(b, "PNG"); return b.getvalue()

def test_delete_moves_out_of_ledger_and_excel_then_restore_puts_it_back(admin):
    admin.post("/api/settings/practice")
    a = sent(admin, customer="DELETE ME CO", amount="321.00")
    b = sent(admin, customer="KEEP ME CO", amount="45.00")
    r0 = row(admin, a)["row"]
    res = admin.post(f"/api/ledger/{a}/delete", json={"reason": "typed twice"}).json()
    assert res["excel"] == "synced"
    assert row(admin, a) is None and row(admin, b)
    gone = admin.get("/api/ledger/deleted").json()
    assert gone[0]["id"] == a and gone[0]["reason"] == "typed twice" and gone[0]["deleted_by"] == "admin"
    ws, lay = mvc(admin)
    assert ws.cell(r0, lay.cols["customer"]).value is None                     # emptied, no rows shifted
    assert ws.cell(row(admin, b)["row"], lay.cols["customer"]).value == "KEEP ME CO"
    an = admin.get("/api/analytics", params={"start": "2026-11-12", "end": "2026-11-12", "by": "customer"}).json()
    assert "DELETE ME CO" not in [x["key"] for x in an.get("breakdown", [])]
    # restore: back in the ledger and written into the workbook again
    assert admin.post(f"/api/ledger/{a}/restore").json()["excel"] == "synced"
    back = row(admin, a)
    ws, lay = mvc(admin)
    assert back and back["row"] and ws.cell(back["row"], lay.cols["customer"]).value == "DELETE ME CO"
    assert admin.get("/api/ledger/deleted").json() == []

def test_delete_and_restore_while_workbook_is_open(admin, app_module):
    import os
    admin.post("/api/settings/practice")
    a = sent(admin, customer="OPEN BOOK CO")
    r0 = row(admin, a)["row"]
    path = admin.get("/api/settings").json()["workbook"]
    lock = os.path.join(os.path.dirname(path), "~$" + os.path.basename(path))
    open(lock, "w").close()
    try:
        assert admin.post(f"/api/ledger/{a}/delete", json={}).json()["excel"] == "pending"
        assert admin.post(f"/api/ledger/{a}/restore").json()["excel"] == "synced"     # removal never ran: just cancelled
    finally:
        os.remove(lock)
    app_module.push_to_excel()
    ws, lay = mvc(admin)
    assert row(admin, a)["row"] == r0 and ws.cell(r0, lay.cols["customer"]).value == "OPEN BOOK CO"

def test_excel_export(admin):
    admin.post("/api/settings/practice")
    a = sent(admin, customer="EXPORT CO", amount="1,000.50", deposit_bank="TNB", taxable="Yes")
    x = admin.post("/api/ledger/export", json={"ids": [a]})
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Ledger", "Daily bank totals"]
    led = wb["Ledger"]
    assert led["B2"].value == "EXPORT CO" and led["H2"].value == 1000.5 and led["A3"].value == "Total"
    t = wb["Daily bank totals"]
    assert [c.value for c in t[2]][:3] == ["2026-11-12", "TNB", 1] and t["H2"].value == 1000.5

def test_page_urls_are_versioned_and_thumbs_trim_the_page(admin):
    """Image URLs carry the stored file's name (a reused page number never shows an old cached picture); the list
    thumbnail trims a check printed on the top of a letter page down to the check."""
    admin.post("/api/settings/practice")
    pid = sent(admin, customer="THUMB CO")
    page = Image.new("RGB", (850, 1100), "white")
    from PIL import ImageDraw
    ImageDraw.Draw(page).rectangle((40, 40, 810, 360), outline="black", width=6)          # a check on the top third
    b = io.BytesIO(); page.save(b, "JPEG")
    p = admin.post(f"/api/payments/{pid}/attachments", files={"files": ("check.jpg", b.getvalue(), "image/jpeg")}).json()["pages"][0]
    assert "?v=" in p["image"] and p["thumb"].startswith(f"/api/pages/{p['id']}/thumb?v=")
    led = row(admin, pid)["pages"][0]
    assert led["image"] == p["image"] and led["thumb"] == p["thumb"]
    t = admin.get(p["thumb"])
    assert t.status_code == 200
    im = Image.open(io.BytesIO(t.content))
    assert im.width <= 480 and im.width / im.height > 2                                  # the check, not the whole page
    assert admin.get(p["image"] + "&download=1").headers["content-disposition"].startswith("attachment")

def test_split_payment_by_invoice(admin, client):
    """One 10,000 check paying three invoices, not all taxable: AR splits it; totals, workbook and exports follow."""
    from fastapi.testclient import TestClient
    admin.post("/api/settings/practice")
    pid = sent(admin, customer="SPLIT CO", amount="10000.00", received="2026-11-13", deposit_bank="TNB")
    email = f"ar.split.{pid}@example.com"
    admin.post("/api/users", json={"email": email, "role": "ar", "password": "pw"})
    ar = TestClient(client.app); ar.post("/api/login", json={"username": email, "password": "pw"})
    url = f"/api/ledger/{pid}/invoices"
    lines = [dict(invoice="inv-1", amount="2000", taxable="Yes"), dict(invoice="INV-2", amount="4,000.00", taxable="No"),
             dict(invoice="INV-3", amount="4000", taxable="Yes")]
    assert ar.post(url, json={"lines": lines[:2]}).status_code == 400                        # adds up to 6,000, not 10,000
    assert "add up to $6,000.00" in ar.post(url, json={"lines": lines[:2]}).json()["detail"]
    assert ar.post(url, json={"lines": lines[:1] + [dict(lines[0], amount="8000")]}).status_code == 400   # same invoice twice
    r = ar.post(url, json={"lines": lines}).json()
    assert r["excel"] == "synced"
    f = row(admin, pid)["fields"]
    assert f["taxable"] == "Partial" and f["taxable_amount"] == "6000.00" and f["invoice"] == "INV-1, INV-2, INV-3"
    assert [l["invoice"] for l in f["invoices"]] == ["INV-1", "INV-2", "INV-3"]
    ws, lay = mvc(admin)
    assert ws.cell(row(admin, pid)["row"], lay.cols["taxable"]).value == "PARTIAL 6,000.00"
    # the exports carry the taxable part
    led = openpyxl.load_workbook(io.BytesIO(admin.post("/api/ledger/export", json={"ids": [pid]}).content))
    assert [led["Ledger"].cell(2, c).value for c in (8, 9, 10, 11)] == [10000, "Partial", 6000, 4000]
    assert "INV-2 4,000.00 not taxable" in led["Ledger"]["P2"].value
    day = [c.value for c in led["Daily bank totals"][2]]
    assert day[7:10] == [6000, 4000, 0]
    # read-only can't split; a plain Yes removes the split
    ro_email = f"ro.split.{pid}@example.com"
    admin.post("/api/users", json={"email": ro_email, "role": "readonly", "password": "pw"})
    ro = TestClient(client.app); ro.post("/api/login", json={"username": ro_email, "password": "pw"})
    assert ro.post(url, json={"lines": lines}).status_code == 403
    assert ar.patch(f"/api/ledger/{pid}", json={"taxable": "Yes"}).status_code == 200
    f = row(admin, pid)["fields"]
    assert f["taxable"] == "Yes" and not f["invoices"] and not f["taxable_amount"]
    ws, lay = mvc(admin)
    assert ws.cell(row(admin, pid)["row"], lay.cols["taxable"]).value == "YES"
    # split again, then remove it with an empty list
    assert ar.post(url, json={"lines": lines}).status_code == 200
    assert ar.post(url, json={"lines": []}).status_code == 200 and row(admin, pid)["fields"]["taxable"] == ""

def test_affects_daily_total(admin, client):
    """On for every payment by default; a negative BPS amount is off unless someone turns it on; anyone's choice wins.
    Payments left out don't count in the daily bank totals sheet."""
    from checkflow.app import in_daily_total
    assert in_daily_total({"amount": "-40", "deposit_bank": "BPS"}) is False
    assert in_daily_total({"amount": "-40", "deposit_bank": "TNB"}) and in_daily_total({"amount": "40", "deposit_bank": "BPS"})
    assert in_daily_total({"amount": "-40", "deposit_bank": "BPS", "daily_total": "Yes"})
    assert not in_daily_total({"amount": "500", "deposit_bank": "TNB", "daily_total": "No"})
    admin.post("/api/settings/practice")
    a = sent(admin, customer="DAILY IN CO", amount="500.00", received="2026-11-14", deposit_bank="BPS")
    b = sent(admin, customer="WITHDRAWAL", amount="-120.00", received="2026-11-14", deposit_bank="BPS", method="cash")
    day = lambda: [c.value for c in openpyxl.load_workbook(io.BytesIO(admin.post("/api/ledger/export", json={"ids": [a, b]}).content))["Daily bank totals"][2]]
    assert day()[2] == 1 and day()[6] == 500                                            # the BPS withdrawal is left out
    assert admin.patch(f"/api/ledger/{b}", json={"daily_total": "Yes"}).json()["excel"] == "unchanged"
    assert day()[2] == 2 and day()[6] == 380
    assert admin.patch(f"/api/ledger/{b}", json={"daily_total": "maybe"}).status_code == 400
    # AR can't change it (it's the cashier's daily total)
    admin.post("/api/users", json={"email": "daily.ar@example.com", "role": "ar", "password": "pw"})
    from fastapi.testclient import TestClient
    ar = TestClient(client.app); ar.post("/api/login", json={"username": "daily.ar@example.com", "password": "pw"})
    assert ar.patch(f"/api/ledger/{b}", json={"daily_total": "No"}).status_code == 403
