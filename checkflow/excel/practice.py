"""A practice workbook laid out like the 2026 payee sheets (one header per day, TAXABLE column, CREDIT then CHECK
subtotals, TOTAL), so the whole flow can be tried without touching the real file. Also used by the tests."""
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

SHEETS = {
    "MVC": ("MAGIC VALLEY CONCRETE 2026", ["DATE", "CUSTOMER", "TICKET REF", "PAYMENT REF:", "TAXABLE", "CASH", "CHECK", "CREDIT",
                                           "DEPOSIT DATE", "DEPOSIT AMOUNT", "BANK"]),
    "UVM": ("UPPER VALLEY MATERIALS 2026", ["DATE", "CUSTOMER", "ORDER #", "TICKET REF", "PAYMENT REF:", "TAXABLE", "CASH", "CHECK",
                                               "CREDIT", "DEPOSIT DATE", "DEPOSIT AMOUNT", "BANK"]),
    "RVP": ("RIO VALLEY PIPE 2026", ["DATE", "CUSTOMER", "ORDER #", "TICKET REF", "PAYMENT REF:", "TAXABLE", "CASH", "CHECK", "CREDIT",
                                     "DEPOSIT DATE", "DEPOSIT AMOUNT", "BANK"]),
    "VCP": ("VALLEY CONCRETE PUMPS 2026", ["DATE", "CUSTOMER", "TICKET REF", "PAYMENT REF:", "TAXABLE", "CASH", "CHECK", "CREDIT",
                                           "DEPOSIT DATE", "DEPOSIT AMOUNT", "BANK"]),
}
SAMPLE = {  # one earlier day per sheet: (customer, ticket, payref, group, amount, bank)
    "MVC": [("PENNY CONSTRUCTION", 35290, 7801, "CREDIT", 1250.00, ""), ("TOVAR & CRUZ", None, 1238, "CHECK", 30666.83, "TNB"),
            ("THOGA INDUSTRIES", None, 1191, "CHECK", 15576.00, "TNB")],
    "UVM": [("ELITE UTILITY & PAVING LLC", 35300, 54, "CHECK", 1200.00, "TNB")],
    "RVP": [("JR GARCIA COMPANY", 35310, 1622, "CHECK", 900.00, "TNB")],
    "VCP": [("MARTINEZ PAVING", None, 44026, "CHECK", 6055.50, "TNB")],
}
GREEN, YELLOW, GREY, PINK, HEAD = "E2EFDA", "FFF2CC", "EDEDED", "FCE4EC", "D9E1F2"

def make_practice(path, day="OCT/08/26"):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    thin = Side(style="thin", color="808080")
    box = Border(top=thin, bottom=thin, left=thin, right=thin)
    for name, (title, hdr) in SHEETS.items():
        ws = wb.create_sheet(name)
        ws.cell(1, 2, title).font = Font(size=16, bold=True)
        col = {h: i + 1 for i, h in enumerate(hdr)}
        def style_row(r, header=False):
            for i, h in enumerate(hdr, 1):
                c = ws.cell(r, i); c.border = box
                if header:
                    c.font = Font(bold=True); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    c.fill = PatternFill("solid", fgColor="404040" if h == "BANK" else HEAD)
                    if h == "BANK": c.font = Font(bold=True, color="FFFFFF")
                else:
                    fill = {"CASH": GREEN, "CHECK": YELLOW, "CREDIT": GREY, "TAXABLE": PINK}.get(h)
                    if fill: c.fill = PatternFill("solid", fgColor=fill)
                    if h in ("CASH", "CHECK", "CREDIT", "DEPOSIT AMOUNT"): c.number_format = "#,##0.00"
        r = 3
        for i, h in enumerate(hdr, 1): ws.cell(r, i, h)
        style_row(r, header=True); ws.row_dimensions[r].height = 30
        subs = []
        for group in ("CREDIT", "CHECK"):
            rows = [x for x in SAMPLE[name] if x[3] == group]
            if not rows: continue
            first = r + 1
            for cust, ticket, payref, g, amt, bank in rows:
                r += 1; style_row(r)
                ws.cell(r, col["DATE"], day); ws.cell(r, col["CUSTOMER"], cust)
                if ticket: ws.cell(r, col["TICKET REF"], ticket)
                ws.cell(r, col["PAYMENT REF:"], payref); ws.cell(r, col[g], amt)
                if bank: ws.cell(r, col["BANK"], bank)
            r += 1; style_row(r)
            r += 1; style_row(r)
            letter = openpyxl.utils.get_column_letter(col[g])
            ws.cell(r, col["DEPOSIT DATE"], group); ws.cell(r, col["DEPOSIT AMOUNT"], f"=SUM({letter}{first}:{letter}{r - 1})")
            subs.append(f"{openpyxl.utils.get_column_letter(col['DEPOSIT AMOUNT'])}{r}")
        if len(subs) > 1:
            r += 1; style_row(r); r += 1; style_row(r)
            ws.cell(r, col["DEPOSIT DATE"], "TOTAL"); ws.cell(r, col["DEPOSIT AMOUNT"], "=" + "+".join(subs))
        for i, h in enumerate(hdr, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = {"CUSTOMER": 34, "DATE": 11}.get(h, 12)
    wb.create_sheet("Sheet1")
    wb.save(path)
    return path
