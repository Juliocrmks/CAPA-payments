"""Write payments into the payee workbook as day blocks that look like the ones people type by hand.

A day block (one per payee sheet per *received* date):
    header row                      (copied from the sheet's last header: values, styles, height)
    entries of the first group      (e.g. CREDIT: card / ACH / wire)
    [blank rows per the sheet's habit]
    subtotal row                    label | =SUM(...)
    entries of the next group       (CHECKS), subtotal ...
    [blank] total row               TOTAL | =J10+J17      (when the day has 2+ groups)
If the sheet's last block is already that date (a second batch the same day), the block is rebuilt with
the old rows kept exactly as they were plus the new ones, so the SUM ranges stay correct."""
import os, re, shutil, tempfile, datetime, time
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter as L_
from .layout import Layout, H, last_row, MAXC, group_of_label
from ..config import METHODS

class WorkbookError(Exception): pass

# ---------------------------------------------------------------- file handling
def lock_path(path):
    d, n = os.path.split(path)
    return os.path.join(d, "~$" + n)

def check_lock(path):
    if os.path.exists(lock_path(path)):
        raise WorkbookError("The workbook is open in Excel. Close it (save first) and try again.")

def load(path):
    if not path or not os.path.exists(path): raise WorkbookError("Workbook not found. Set the workbook path in Settings.")
    check_lock(path)
    try:
        return openpyxl.load_workbook(path, keep_vba=path.lower().endswith(".xlsm"))
    except PermissionError:
        raise WorkbookError("The workbook is open in Excel. Close it and try again.")

def backup(path, keep=30, min_age=0):
    d = os.path.join(os.path.dirname(path), "checkflow_backups")
    os.makedirs(d, exist_ok=True)
    if min_age:
        ts = [os.path.getmtime(os.path.join(d, f)) for f in os.listdir(d)]
        if ts and time.time() - max(ts) < min_age: return None
    stem, ext = os.path.splitext(os.path.basename(path))
    b = os.path.join(d, f"{stem}_{datetime.datetime.now():%Y%m%d_%H%M%S}{ext}")
    shutil.copy2(path, b)
    for f in sorted(os.listdir(d))[:-keep]:
        try: os.remove(os.path.join(d, f))
        except OSError: pass
    return b

def save(wb, path, backup_min_age=0):
    check_lock(path)
    backup(path, min_age=backup_min_age)
    fd, tmp = tempfile.mkstemp(suffix=os.path.splitext(path)[1], dir=os.path.dirname(path) or ".")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError:
        raise WorkbookError("The workbook is open in Excel. Close it and try again.")
    finally:
        if os.path.exists(tmp):
            try: os.remove(tmp)
            except OSError: pass

def list_sheets(path, ignore=()):
    wb = openpyxl.load_workbook(path, read_only=True)
    try: names = list(wb.sheetnames)
    finally: wb.close()
    ig = {H(x) for x in ignore}
    return [n for n in names if H(n) not in ig]

# ---------------------------------------------------------------- values
def date_text(iso):
    return datetime.date.fromisoformat(iso).strftime("%b/%d/%y").upper()     # OCT/09/26

def num_or_text(v):
    v = str(v or "").strip()
    if re.fullmatch(r"\d{1,12}", v) and not (len(v) > 1 and v.startswith("0") and len(v) > 6):
        return int(v)
    return v or None

def group_for(method): return METHODS.get(method, {}).get("column", "check")

# ---------------------------------------------------------------- styles
def _row_snapshot(ws, r):
    return dict(height=ws.row_dimensions[r].height,
                cells={c: (ws.cell(r, c).value, copy(ws.cell(r, c)._style), ws.cell(r, c).number_format) for c in range(1, MAXC)})

def _paint(ws, r, snap, values=True):
    if snap.get("height"): ws.row_dimensions[r].height = snap["height"]
    for c, (v, st, nf) in snap["cells"].items():
        cell = ws.cell(r, c)
        cell._style = copy(st)
        if values:
            cell.value = v
            if nf: cell.number_format = nf

def _plain(cell):
    f = cell.fill
    return not (f and f.fill_type) or str(getattr(f.fgColor, "rgb", "")) in ("00000000", "FFFFFFFF")

def _templates(lay):
    """Representative rows to copy styles from: one 'normal' entry row per group, one subtotal, one total, one blank."""
    ws = lay.ws
    lo = max(2, last_row(ws) - 600)
    entry, sub, total = {}, None, None
    for r in range(last_row(ws), lo, -1):
        if lay.is_entry(r):
            g = lay.entry_group(r)
            if g and g not in entry and _plain(ws.cell(r, lay.cols["customer"])): entry[g] = r
        v = ws.cell(r, lay.sum_col).value
        if isinstance(v, str) and v.startswith("="):
            lab = ws.cell(r, lay.label_col).value
            if group_of_label(lab) and sub is None: sub = r
            elif "TOTAL" in H(lab) and total is None: total = r
    any_entry = next(iter(entry.values()), None)
    blank_row = last_row(ws) + 3
    return dict(entry={g: _row_snapshot(ws, entry.get(g, any_entry)) for g in ("cash", "check", "credit")} if any_entry else {},
                sub=_row_snapshot(ws, sub) if sub else None, total=_row_snapshot(ws, total or sub) if (total or sub) else None,
                blank=_row_snapshot(ws, blank_row))

def _clear(ws, r, blank):
    if blank: _paint(ws, r, blank, values=True)
    else:
        for c in range(1, MAXC): ws.cell(r, c).value = None

# ---------------------------------------------------------------- write a day
def _entry_snapshot(lay, tpl, it):
    g = group_for(it["method"])
    base = tpl["entry"].get(g) or tpl["entry"].get("check")
    snap = dict(height=base["height"] if base else None, cells={c: (None, st, nf) for c, (v, st, nf) in (base or {"cells": {}})["cells"].items()})
    cols, vals = lay.cols, {}
    vals[cols["date"]] = date_text(it["received"])
    vals[cols["customer"]] = (it.get("customer") or "").upper()
    if "order" in cols and it.get("order"): vals[cols["order"]] = it["order"].upper()
    ticket = it.get("ticket") or ("WIRE" if it["method"] == "wire" else None)
    if "ticket" in cols and ticket: vals[cols["ticket"]] = num_or_text(ticket)
    if "payref" in cols:
        if it["method"] == "wire" and not it.get("payref"):
            vals[cols["payref"]] = datetime.datetime.fromisoformat(it["received"])
        elif it.get("payref"):
            vals[cols["payref"]] = num_or_text(it["payref"])
    if "taxable" in cols and it.get("taxable") in ("Yes", "No"): vals[cols["taxable"]] = it["taxable"].upper()
    col = lay.amount_cols.get(g) or lay.amount_cols.get("credit") or lay.cols["check"]
    vals[col] = round(float(it["amount"]), 2)
    if "bank" in cols and it.get("deposit_bank"): vals[cols["bank"]] = it["deposit_bank"].upper()
    for c, v in vals.items():
        st = snap["cells"].get(c, (None, copy(openpyxl.cell.cell.Cell(lay.ws)._style), "General"))
        nf = ("d-mmm" if isinstance(v, datetime.datetime) else "#,##0.00" if c == col else
              "General" if isinstance(v, int) or c in (cols.get("ticket"), cols.get("payref")) else st[2])
        snap["cells"][c] = (v, st[1], nf)
    return snap, g, vals

def write_day(ws, received_iso, items):
    """Write items (same sheet, same received date) into ws. -> {item id: row}"""
    lay = Layout(ws)
    tpl = _templates(lay)
    dt = date_text(received_iso)
    groups = {g: [] for g in ("cash", "check", "credit")}
    # 1) where does the block go?
    h = lay.hdr
    _, end = lay.block_bounds(h)
    existing_date = lay.block_date(h)
    entries = [r for r in range(h + 1, end + 1) if lay.is_entry(r)]
    empty_header = not entries
    same_day = existing_date == dt and all(H(ws.cell(r, lay.cols["date"]).value) in (dt, "") for r in entries)
    if same_day or empty_header:
        for r in range(h + 1, end + 1):           # keep the rows already typed for this day, exactly as they are
            if lay.is_entry(r): groups[lay.entry_group(r)].append(("old", _row_snapshot(ws, r)))
        for r in range(h + 1, end + 1): _clear(ws, r, tpl["blank"])
    else:
        hsnap = _row_snapshot(ws, h)
        h = end + 1 + lay.gap_between_blocks
        for r in range(end + 1, h): _clear(ws, r, tpl["blank"])
        _paint(ws, h, hsnap)
    out = {}
    for it in items:
        snap, g, _ = _entry_snapshot(lay, tpl, it)
        groups[g].append((it["id"], snap))
    # 2) rows
    r = h + 1
    subs = []
    sumL = L_(lay.sum_col)
    for g in lay.order:
        rows = groups.get(g) or []
        if not rows: continue
        first = r
        for key, snap in rows:
            _paint(ws, r, snap)
            if key != "old": out[key] = r
            r += 1
        for _ in range(lay.gap_before_subtotal):
            _clear(ws, r, tpl["blank"]); r += 1
        col = L_(lay.amount_cols.get(g) or lay.cols["check"])
        if tpl["sub"]: _paint(ws, r, tpl["sub"], values=False)
        ws.cell(r, lay.label_col).value = lay.labels[g]
        ws.cell(r, lay.sum_col).value = f"=SUM({col}{first}:{col}{r - 1})"
        ws.cell(r, lay.sum_col).number_format = "#,##0.00"
        subs.append(f"{sumL}{r}")
        r += 1
    if len(subs) >= 2:
        _clear(ws, r, tpl["blank"]); r += 1
        if tpl["total"]: _paint(ws, r, tpl["total"], values=False)
        ws.cell(r, lay.label_col).value = lay.total_label
        ws.cell(r, lay.sum_col).value = "=" + "+".join(subs)
        ws.cell(r, lay.sum_col).number_format = "#,##0.00"
        r += 1
    return out

def write_payments(path, items, now_backup=True):
    """items: [{id, sheet, received, customer, order, ticket, payref, method, amount, taxable, deposit_bank}]
    -> {id: (sheet, row)}. All-or-nothing: the file is only replaced if every day block was written."""
    wb = load(path)
    out = {}
    by = {}
    for it in items:
        if it["sheet"] not in wb.sheetnames: raise WorkbookError(f"Sheet '{it['sheet']}' does not exist in the workbook.")
        by.setdefault((it["sheet"], it["received"]), []).append(it)
    for (sheet, day) in sorted(by, key=lambda k: (k[1], k[0])):
        try:
            rows = write_day(wb[sheet], day, by[(sheet, day)])
        except ValueError as e:
            raise WorkbookError(str(e))
        out.update({k: (sheet, v) for k, v in rows.items()})
    save(wb, path)
    return out

# ---------------------------------------------------------------- edits after the fact
def find_row(lay, it, hint=None):
    """Locate a written payment by date + customer + (payment ref or amount). Checks the remembered row first."""
    ws, cols = lay.ws, lay.cols
    dt = date_text(it["received"]) if it.get("received") else None
    want_c = H(it.get("customer"))
    want_amt = round(float(it["amount"]), 2) if it.get("amount") not in (None, "") else None
    def ok(r):
        if H(ws.cell(r, cols["customer"]).value) != want_c: return False
        if dt and H(ws.cell(r, cols["date"]).value) not in (dt, ""): return False
        amts = [ws.cell(r, c).value for c in lay.amount_cols.values()]
        if want_amt is not None and not any(isinstance(a, (int, float)) and abs(a - want_amt) < 0.005 for a in amts): return False
        return True
    if hint and hint <= ws.max_row and ok(hint): return hint
    for r in range(last_row(ws), max(1, last_row(ws) - 5000), -1):
        if ok(r): return r
    return None

def update_rows(path, updates):
    """updates: [{id, sheet, row, received, customer, amount, taxable, deposit_bank}] -> {id: ('ok', row) | ('err', msg)}"""
    wb = load(path)
    out, changed, lays = {}, False, {}
    for u in updates:
        if u["sheet"] not in wb.sheetnames: out[u["id"]] = ("err", f"Sheet '{u['sheet']}' not found."); continue
        lay = lays.get(u["sheet"]) or lays.setdefault(u["sheet"], Layout(wb[u["sheet"]]))
        r = find_row(lay, u, u.get("row"))
        if not r: out[u["id"]] = ("err", "Row not found in the sheet (was it moved or edited by hand?)."); continue
        if "taxable" in lay.cols and "taxable" in u:
            lay.ws.cell(r, lay.cols["taxable"]).value = u["taxable"].upper() if u["taxable"] in ("Yes", "No") else None
        if "bank" in lay.cols and "deposit_bank" in u:
            lay.ws.cell(r, lay.cols["bank"]).value = (u["deposit_bank"] or "").upper() or None
        out[u["id"]] = ("ok", r); changed = True
    if changed: save(wb, path, backup_min_age=600)
    return out
