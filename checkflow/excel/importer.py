"""Read existing rows out of the workbook.

customers(path)  -> distinct customer names per sheet (used to snap OCR'd names to what the sheets already say)
history(path)    -> every entry row as a ledger record (optional one-time migration of past days into the ledger)"""
import datetime, re
import openpyxl
from .layout import Layout, H

MON = {m: i for i, m in enumerate(["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"], 1)}

def parse_sheet_date(v):
    if isinstance(v, (datetime.datetime, datetime.date)): return (v.date() if isinstance(v, datetime.datetime) else v).isoformat()
    m = re.match(r"^\s*([A-Za-z]{3})[/_ ]?(\d{1,2})/(\d{2,4})\s*$", str(v or ""))
    if not m or m.group(1).upper() not in MON: return None
    y = int(m.group(3)); y = 2000 + y if y < 100 else y
    try: return datetime.date(y, MON[m.group(1).upper()], int(m.group(2))).isoformat()
    except ValueError: return None

def _iter_entries(ws):
    try: lay = Layout(ws)
    except ValueError: return
    hs = lay.headers
    from .layout import header_map
    for k, h in enumerate(hs):
        cols = header_map(ws, h)
        nxt = hs[k + 1] if k + 1 < len(hs) else ws.max_row + 1
        amount_cols = {g: cols[g] for g in ("cash", "check", "credit") if g in cols}
        last_date = None
        for r in range(h + 1, nxt):
            cust = ws.cell(r, cols["customer"]).value
            if cust in (None, "") or H(cust) == "CUSTOMER": continue
            g = next((g for g, c in amount_cols.items() if isinstance(ws.cell(r, c).value, (int, float))), None)
            if not g: continue
            d = parse_sheet_date(ws.cell(r, cols["date"]).value) or last_date
            last_date = d
            yield r, cols, g, d, amount_cols

def customers(path, sheets=None):
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    out = {}
    try:
        for ws in wb.worksheets:
            if sheets and ws.title not in sheets: continue
            for r, cols, g, d, _ in _iter_entries(ws):
                name = re.sub(r"\s+", " ", str(ws.cell(r, cols["customer"]).value)).strip().upper()
                if len(name) >= 2 and not name.startswith(("WITHDRAWAL", "DEPOSIT")): out.setdefault(name, set()).add(ws.title)
    finally:
        wb.close()
    return {k: sorted(v) for k, v in out.items()}

def history(path, sheets=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    method_of = {"cash": "cash", "check": "check", "credit": "card"}
    try:
        for ws in wb.worksheets:
            if sheets and ws.title not in sheets: continue
            for r, cols, g, d, amount_cols in _iter_entries(ws):
                def val(k):
                    c = cols.get(k)
                    v = ws.cell(r, c).value if c else None
                    if isinstance(v, datetime.datetime): return v.strftime("%-d-%b") if k == "payref" else v.date().isoformat()
                    if isinstance(v, float) and v.is_integer(): v = int(v)
                    return "" if v is None else str(v).strip()
                ticket = val("ticket")
                method = method_of[g]
                if g == "credit" and ticket.upper() == "WIRE": method, ticket = "wire", ""
                tax = val("taxable").upper()
                yield dict(sheet=ws.title, row=r, received=d or "", customer=val("customer").upper(), order=val("order"),
                           ticket=ticket, payref=val("payref"), method=method, amount=f"{float(ws.cell(r, amount_cols[g]).value):.2f}",
                           taxable={"YES": "Yes", "NO": "No"}.get(tax, ""), deposit_bank=val("bank").upper())
    finally:
        wb.close()
