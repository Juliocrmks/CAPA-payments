"""CheckFlow API + static UI. Run with:  python run.py   (or CheckFlow.command / run.bat)"""
import os, json, threading, queue, datetime, uuid, traceback
import cv2
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request, Response, Depends, Query
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from typing import Optional
from . import db, auth, config
from .auth import require_user, require_admin, require_perm, can
from .excel import writer, importer

app = FastAPI(title="CheckFlow")
WB_LOCK = threading.Lock()

def startup():
    db.init()
    auth.seed_admin()
startup()

def workbook_path():
    p = db.get_setting("workbook", "")
    if not p:
        p = os.path.join(config.DATA, "Practice_Workbook.xlsx")
        if not os.path.exists(p):
            from .excel.practice import make_practice
            make_practice(p)
        db.set_setting("workbook", p)
    return p

def book_sheet(name, book_sheets, payees=None):
    """The workbook's own sheet for a payee sheet name, through the payee aliases (UVM = UPPER, RVP = RIO), so a payment
    marked UVM still goes to a workbook whose sheet is called UPPER and the other way round."""
    if not name: return None
    up = {s.upper(): s for s in book_sheets}
    if name.upper() in up: return up[name.upper()]
    for _n, sheet, aliases in (payees if payees is not None else db.registry().payees()):
        names = [sheet.upper()] + [a.upper() for a in aliases]
        if name.upper() in names:
            for n in names:
                if n in up: return up[n]
    return None

def ignore_list(): return [s.strip() for s in (db.get_setting("ignore_sheets") or "").split(",") if s.strip()]

def sheets():
    try: return [s for s in writer.list_sheets(workbook_path(), ignore_list()) if s]
    except Exception: return []

# ================================================================== auth
class Cred(BaseModel):
    username: str
    password: str

@app.get("/api/me")
def me(request: Request):
    return dict(user=auth.current_user(request), auth_mode=db.get_setting("auth_mode"))

@app.post("/api/login")
def login(c: Cred, request: Request, response: Response):
    tok = auth.login(c.username, c.password, request.client.host if request.client else "")
    response.set_cookie(auth.COOKIE, tok, httponly=True, samesite="lax", max_age=auth.TTL)
    return {"ok": True}

@app.post("/api/logout")
def logout(request: Request, response: Response):
    auth.logout(request.cookies.get(auth.COOKIE, "")); response.delete_cookie(auth.COOKIE); return {"ok": True}

class PwChange(BaseModel):
    current: str = ""
    new: str

@app.post("/api/me/password")
def change_password(p: PwChange, user=Depends(require_user)):
    """Anyone changes their own password, to anything they like (required at first sign-in / after an admin reset)."""
    if user["source"] != "local": raise HTTPException(400, "Your password is managed by Active Directory.")
    if not auth._local_check(user["username"], p.current): raise HTTPException(400, "Current password is wrong.")
    if not p.new: raise HTTPException(400, "Enter a new password.")
    if user["must_change"] and p.new == p.current: raise HTTPException(400, "Choose a password different from the temporary one.")
    auth.set_password(user["id"], p.new)
    db.audit(user["username"], "password_changed")
    return {"ok": True}

class ResetAsk(BaseModel):
    login: str

_asks = {}
@app.post("/api/password-reset-request")
def password_reset_request(r: ResetAsk, request: Request):
    """'I forgot my password': admins see the request in Settings > Users and set a temporary password.
    Same answer whether or not the account exists (no way to probe for accounts)."""
    import time
    ip = request.client.host if request.client else ""
    now = time.time(); _asks[ip] = [t for t in _asks.get(ip, []) if now - t < 3600]
    if len(_asks[ip]) < 10:
        _asks[ip].append(now)
        u = auth.find_login(r.login)
        if u and u["source"] == "local" and u["active"]:
            db.add_reset_request(r.login.strip().lower(), u["id"]); db.audit(u["username"], "password_reset_requested")
    return dict(ok=True, message="Request sent. An administrator will set a temporary password and give it to you.")

@app.get("/api/password-requests")
def password_requests(_=Depends(require_admin)): return db.reset_requests()

@app.post("/api/password-requests/{rid}/dismiss")
def dismiss_request(rid: int, me_=Depends(require_admin)):
    db.close_reset_requests(req_id=rid, by=me_["username"]); return db.reset_requests()

def _users():
    with db.conn() as c:
        return [dict(r) for r in c.execute("""SELECT id, username, email, role, source, display_name, active, must_change, last_login, created_at
                                              FROM users ORDER BY active DESC, CASE role WHEN 'admin' THEN 0 WHEN 'ar' THEN 1 ELSE 2 END, display_name, username""")]

@app.get("/api/users")
def users(_=Depends(require_admin)): return _users()

@app.get("/api/roles")
def roles(_=Depends(require_admin)): return auth.ROLES

class NewUser(BaseModel):
    email: str = ""
    username: str = ""
    display_name: str = ""
    role: str = "ar"
    password: str = ""

@app.post("/api/users")
def add_user(u: NewUser, me_=Depends(require_admin)):
    """No password given -> a temporary one is generated, shown once, and the user must choose their own at first sign-in."""
    pw = auth.create_user(u.username, u.password, u.role, display_name=u.display_name, email=u.email, must_change=True)
    db.audit(me_["username"], "user_added", target=u.email or u.username, role=u.role)
    return dict(users=_users(), password=pw)

def _role_of(text):
    t = (text or "").lower()
    if "admin" in t: return "admin"
    if "read" in t or "lectura" in t or "only" in t or "consulta" in t: return "readonly"
    return "ar"

class BulkUsers(BaseModel):
    text: str

@app.post("/api/users/bulk")
def add_users_bulk(b: BulkUsers, me_=Depends(require_admin)):
    """One person per line: Name, email, role   (tab or comma separated; role words: admin / AR / read only)."""
    import re as _re
    out = []
    for line in b.text.splitlines():
        if not line.strip(): continue
        m = _re.search(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", line)
        email = m.group(0) if m else ""
        rest = _re.sub(r"\[?[\w.+-]+@[\w-]+(?:\.[\w-]+)+\]?(\(mailto:[^)]*\))?", "\t", line)
        parts = [x.strip() for x in _re.split(r"\t|,|;|\s{2,}", rest) if x.strip()]
        role = "ar"
        if len(parts) >= 2: role = _role_of(parts[-1]); parts = parts[:-1]
        name = " ".join(parts)
        try:
            pw = auth.create_user("", "", role, display_name=name, email=email)
            out.append(dict(name=name, email=email.lower(), role=role, ok=True, password=pw))
            db.audit(me_["username"], "user_added", target=email, role=role)
        except HTTPException as e:
            out.append(dict(name=name, email=email.lower(), role=role, ok=False, error=e.detail))
    return dict(results=out, users=_users())

class UserPatch(BaseModel):
    active: bool | None = None
    role: str | None = None
    display_name: str | None = None
    email: str | None = None

def _admins_left(c, excluding):
    return c.execute("SELECT COUNT(*) n FROM users WHERE role='admin' AND active=1 AND id<>?", (excluding,)).fetchone()["n"]

@app.patch("/api/users/{uid}")
def edit_user(uid: int, p: UserPatch, me_=Depends(require_admin)):
    with db._lock, db.conn() as c:
        if not c.execute("SELECT 1 FROM users WHERE id=?", (uid,)).fetchone(): raise HTTPException(404)
        if p.active is not None:
            if uid == me_["id"] and not p.active: raise HTTPException(400, "You can't disable your own account.")
            if not p.active and not _admins_left(c, uid): raise HTTPException(400, "Keep at least one active admin.")
            c.execute("UPDATE users SET active=? WHERE id=?", (int(p.active), uid))
            if not p.active: c.execute("DELETE FROM sessions WHERE user_id=?", (uid,))
        if p.role is not None:
            if p.role not in auth.ROLES: raise HTTPException(400, "Unknown role.")
            if uid == me_["id"] and p.role != "admin": raise HTTPException(400, "You can't remove your own admin role.")
            if p.role != "admin" and not _admins_left(c, uid): raise HTTPException(400, "Keep at least one active admin.")
            c.execute("UPDATE users SET role=? WHERE id=?", (p.role, uid))
        if p.display_name is not None: c.execute("UPDATE users SET display_name=? WHERE id=?", (p.display_name.strip(), uid))
        if p.email is not None:
            em = p.email.strip().lower()
            if em and c.execute("SELECT 1 FROM users WHERE id<>? AND (email=? COLLATE NOCASE OR username=? COLLATE NOCASE)", (uid, em, em)).fetchone():
                raise HTTPException(400, f"{em} already has an account.")
            c.execute("UPDATE users SET email=? WHERE id=?", (em, uid))
    db.audit(me_["username"], "user_changed", target=uid, **p.model_dump(exclude_none=True))
    return _users()

class PwReset(BaseModel):
    password: str = ""
    must_change: bool = True

@app.post("/api/users/{uid}/password")
def reset_password(uid: int, r: PwReset, me_=Depends(require_admin)):
    """Admin sets a password for someone who lost theirs: typed, or generated when left empty. By default the person
    must choose their own at the next sign-in. Their open sessions end."""
    with db.conn() as c:
        u = c.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
    if not u: raise HTTPException(404)
    if u["source"] != "local": raise HTTPException(400, "This password is managed by Active Directory.")
    pw = r.password or auth.temp_password()
    auth.set_password(uid, pw, must_change=r.must_change)
    with db._lock, db.conn() as c:
        if uid != me_["id"]: c.execute("DELETE FROM sessions WHERE user_id=?", (uid,))
    db.close_reset_requests(user_id=uid, by=me_["username"])
    db.audit(me_["username"], "password_reset", target=u["username"], must_change=r.must_change)
    return dict(password=pw, users=_users())

# ================================================================== settings
SETTING_KEYS = ["workbook", "ignore_sheets", "deposit_banks", "deposit_bank_check", "deposit_bank_cash", "deposit_bank_card",
                "deposit_bank_ach", "deposit_bank_wire", "keep_sensitive_pages", "auth_mode", "ldap_server", "ldap_domain",
                "ldap_upn_suffix", "ldap_base_dn", "ldap_admin_group", "ldap_ar_group"]

@app.get("/api/settings")
def get_settings(_=Depends(require_admin)):
    s = db.settings(); p = workbook_path()
    return dict({k: s.get(k, "") for k in SETTING_KEYS}, workbook=p, workbook_exists=os.path.exists(p),
                workbook_open=os.path.exists(writer.lock_path(p)), sheets=sheets(), methods=config.METHODS)

@app.post("/api/settings")
def save_settings(s: dict, me_=Depends(require_admin)):
    for k, v in s.items():
        if k not in SETTING_KEYS: continue
        if k == "workbook":
            v = (v or "").strip().strip('"').strip("'")
            if v and not os.path.exists(v): raise HTTPException(400, "That workbook file doesn't exist. Check the path.")
            if v and not v.lower().endswith((".xlsx", ".xlsm")): raise HTTPException(400, "Choose an .xlsx or .xlsm workbook.")
        if k == "auth_mode" and v not in ("local", "ldap"): raise HTTPException(400, "Sign-in must be local or ldap.")
        db.set_setting(k, v)
    db.audit(me_["username"], "settings_changed", keys=[k for k in s if k in SETTING_KEYS])
    return get_settings()

@app.post("/api/settings/practice")
def practice(me_=Depends(require_admin)):
    from .excel.practice import make_practice
    p = os.path.join(config.DATA, "Practice_Workbook.xlsx"); make_practice(p); db.set_setting("workbook", p)
    return get_settings()

# ------------------------------------------------------------------ workbooks
# Which workbooks are there to pick from: CheckFlow's own folder (created / uploaded), the practice one, the folder of
# the workbook in use (e.g. next to the real file on a Mac) and every workbook payments were written to.
import re as _re
XL = (".xlsx", ".xlsm")

def _book_dirs():
    return list(dict.fromkeys([config.WORKBOOKS, config.DATA, os.path.dirname(workbook_path())]))

def _books():
    cur, counts, seen, out = workbook_path(), db.workbook_counts(), set(), []
    paths = []
    for d in _book_dirs():
        try: paths += [os.path.join(d, n) for n in sorted(os.listdir(d))]
        except OSError: pass
    paths += list(counts) + [cur]
    for p in paths:
        if not p or not os.path.isabs(p): continue      # e.g. 'upload:...' tags of imported rows
        p = os.path.normpath(p); n = os.path.basename(p)
        if p in seen or not n.lower().endswith(XL) or n.startswith(("~$", ".")): continue
        seen.add(p)
        ex = os.path.isfile(p); st = os.stat(p) if ex else None
        out.append(dict(path=p, name=n, folder=os.path.dirname(p), exists=ex, size=st.st_size if st else 0,
                        modified=datetime.datetime.fromtimestamp(st.st_mtime).isoformat(timespec="minutes") if st else "",
                        current=p == os.path.normpath(cur), open=os.path.exists(writer.lock_path(p)), payments=counts.get(p, 0),
                        practice=n == "Practice_Workbook.xlsx" and os.path.dirname(p) == os.path.normpath(config.DATA)))
    out.sort(key=lambda b: (not b["current"], b["practice"], -(datetime.datetime.fromisoformat(b["modified"]).timestamp() if b["modified"] else 0)))
    return out

def _listed(path):
    p = os.path.normpath(path) if path and os.path.isabs(path) else ""
    b = next((b for b in _books() if p and b["path"] == p), None)
    if not b: raise HTTPException(404, "That workbook isn't in the list.")
    return b

def _new_book_dir():
    d = os.path.dirname(os.path.normpath(workbook_path()))
    return config.WORKBOOKS if d == os.path.normpath(config.DATA) else d

def _clean_name(name):
    n = _re.sub(r'[\\/:*?"<>|\x00-\x1f]', "", (name or "")).strip().strip(".")
    if not n: raise HTTPException(400, "Give the workbook a name.")
    return n if n.lower().endswith(XL) else n + ".xlsx"

@app.get("/api/workbooks")
def list_workbooks(_=Depends(require_admin)):
    return dict(workbooks=_books(), new_folder=_new_book_dir())

class PickBook(BaseModel):
    path: str

@app.post("/api/workbooks/select")
def select_workbook(b: PickBook, me_=Depends(require_admin)):
    w = _listed(b.path)
    if not w["exists"]: raise HTTPException(400, "That workbook file isn't there anymore.")
    db.set_setting("workbook", w["path"]); db.audit(me_["username"], "workbook_selected", file=w["name"])
    return dict(settings=get_settings(), workbooks=_books())

class NewBook(BaseModel):
    name: str
    year: Optional[int] = None
    source: str = "current"          # current: copy the layout of the workbook in use | standard: CAPA's standard layout
    use: bool = True

@app.post("/api/workbooks/new")
def new_workbook(b: NewBook, me_=Depends(require_admin)):
    name = _clean_name(b.name)
    folder = _new_book_dir(); os.makedirs(folder, exist_ok=True)
    dst = os.path.join(folder, name)
    if os.path.exists(dst): raise HTTPException(409, f"There's already a workbook called {name} in that folder.")
    if b.year is not None and not 1990 <= b.year <= 2100: raise HTTPException(400, "Year looks wrong.")
    if b.source == "standard":
        from .excel.practice import make_practice
        import tempfile
        src = os.path.join(tempfile.mkdtemp(), "standard.xlsx"); make_practice(src); ignore = ["Sheet1"]
    else:
        src, ignore = workbook_path(), ignore_list()
    try:
        with WB_LOCK: res = writer.new_workbook(src, dst, year=b.year, ignore=ignore)
    except writer.WorkbookError as e: raise HTTPException(400, str(e))
    if b.use: db.set_setting("workbook", dst)
    db.audit(me_["username"], "workbook_created", file=name, source=b.source, year=b.year, sheets=res["sheets"])
    return dict(res, path=dst, settings=get_settings(), workbooks=_books())

@app.post("/api/workbooks/upload")
async def upload_workbook(file: UploadFile = File(...), me_=Depends(require_admin)):
    name = _clean_name(file.filename)
    os.makedirs(config.WORKBOOKS, exist_ok=True)
    base, ext = os.path.splitext(name); dst, i = os.path.join(config.WORKBOOKS, name), 2
    while os.path.exists(dst): dst, i = os.path.join(config.WORKBOOKS, f"{base} ({i}){ext}"), i + 1
    data = await file.read()
    if len(data) > 80 * 1024 * 1024: raise HTTPException(400, "That file is over 80 MB.")
    with open(dst, "wb") as f: f.write(data)
    try: sheets_ = writer.list_sheets(dst)
    except Exception:
        os.remove(dst); raise HTTPException(400, "That file isn't an Excel workbook CheckFlow can open (.xlsx or .xlsm).")
    db.audit(me_["username"], "workbook_uploaded", file=os.path.basename(dst))
    return dict(path=dst, sheets=sheets_, workbooks=_books())

@app.get("/api/workbooks/download")
def download_workbook(path: str = Query(...), me_=Depends(require_admin)):
    w = _listed(path)
    if not w["exists"]: raise HTTPException(404, "That workbook file isn't there anymore.")
    db.audit(me_["username"], "workbook_downloaded", file=w["name"])
    return FileResponse(w["path"], filename=w["name"],
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# ------------------------------------------------------------------ sync the workbook from the ledger
# The ledger is the source of truth. Sync rewrites the payee sheets of the workbook in use so they hold exactly the
# ledger's payments for the workbook's year, in the workbook's own format. Rows that are only in the workbook are listed
# first and, by default, added to the ledger so nothing typed by hand is lost.
def _sync_state(year=None):
    from collections import Counter, defaultdict, deque
    from .excel.layout import Layout
    from .excel import ledger_import as li
    path = workbook_path()
    wb = writer.load(path)
    ig = {s.strip().upper() for s in ignore_list()}
    book, recs = [], []
    for ws in wb.worksheets:
        if ws.title.strip().upper() in ig: continue
        try: Layout(ws)
        except ValueError: continue
        book.append(ws.title); recs += list(importer.history_ws(ws))
    undated = [r for r in recs if not r["received"]]
    recs = [r for r in recs if r["received"]]
    li.flag_dates(recs)                  # likely year typos ('JAN/08/25' among 2026 rows): the fix is used only as a 2nd key
    flagged = [r for r in recs if r.get("date_issue")]
    if not year:
        ys = Counter(r["received"][:4] for r in recs)
        m = _re.search(r"(19|20)\d{2}", os.path.basename(path))
        year = int(ys.most_common(1)[0][0]) if ys else int(m.group(0)) if m else datetime.date.today().year
    payees, led, unmapped = db.registry().payees(), [], Counter()
    for x in db.ledger_rows():
        if not (x["fields"].get("received") or "").startswith(str(year)): continue
        s = book_sheet(x["sheet"], book, payees)
        if s: led.append((x, s))
        else: unmapped[x["sheet"] or "(no payee)"] += 1
    fp = lambda s, d, f: li.fingerprint(s, d, f["amount"], f["customer"])
    def fixed_date(r):
        fx = (r.get("date_issue") or {}).get("fix")
        return fx if fx and not r["received"].startswith(str(year)) and fx.startswith(str(year)) else None
    fixed = 0
    q = defaultdict(deque)
    for r in sorted(recs, key=lambda r: (r["sheet"], r["row"])):
        q[fp(r["sheet"], r["received"], r)].append(r)
        fx = (r.get("date_issue") or {}).get("fix")
        if fx: q[fp(r["sheet"], fx, r)].append(r)          # matches a ledger row imported with the corrected date
    used, pos, only_led = set(), {}, []
    for x, s in sorted(led, key=lambda t: t[0]["id"]):
        dq = q[fp(s, x["fields"]["received"], x["fields"])]
        while dq and id(dq[0]) in used: dq.popleft()
        if dq: r = dq.popleft(); used.add(id(r)); pos[x["id"]] = r["row"]
        else: only_led.append((x, s))
    only_wb = []
    for r in sorted(recs, key=lambda r: (r["sheet"], r["row"])):
        if id(r) in used: continue
        fx = fixed_date(r)
        if fx: fixed += 1
        only_wb.append(dict(r, received=fx or r["received"]))
    per = {s: dict(sheet=s, ledger=0, ledger_total=0.0, workbook=0, workbook_total=0.0, add=0, remove=0) for s in book}
    for x, s in led: per[s]["ledger"] += 1; per[s]["ledger_total"] += float(x["fields"]["amount"] or 0)
    for r in recs: per[r["sheet"]]["workbook"] += 1; per[r["sheet"]]["workbook_total"] += float(r["amount"])
    for x, s in only_led: per[s]["add"] += 1
    for r in only_wb: per[r["sheet"]]["remove"] += 1
    for p in per.values(): p["ledger_total"], p["workbook_total"] = round(p["ledger_total"], 2), round(p["workbook_total"], 2)
    return dict(path=path, year=year, book=book, led=led, pos=pos, only_wb=only_wb, only_led=only_led, per=list(per.values()),
                unmapped=dict(unmapped), undated=undated, fixed=fixed, flagged=flagged)

def _row_out(r): return dict(sheet=r["sheet"], row=r["row"], received=r["received"], customer=r["customer"], amount=r["amount"], method=r["method"])

@app.get("/api/workbooks/sync")
def sync_preview(year: Optional[int] = None, _=Depends(require_admin)):
    try: st = _sync_state(year)
    except writer.WorkbookError as e: raise HTTPException(409, str(e))
    other = [r for r in st["only_wb"] if not r["received"].startswith(str(st["year"]))]
    return dict(workbook=os.path.basename(st["path"]), year=st["year"], sheets=st["per"],
                ledger=len(st["led"]), add=len(st["only_led"]), remove=len(st["only_wb"]), other_year=len(other),
                only_in_workbook=[_row_out(r) for r in st["only_wb"][:200]], unmapped=st["unmapped"],
                undated=[dict(sheet=r["sheet"], row=r["row"], customer=r["customer"], amount=r["amount"]) for r in st["undated"]],
                dates_fixed=st["fixed"], in_sync=not st["only_led"] and not st["only_wb"])

class SyncReq(BaseModel):
    year: Optional[int] = None
    keep_workbook_rows: bool = True      # add rows that are only in the workbook to the ledger first

@app.post("/api/workbooks/sync")
def sync_workbook(r: SyncReq, me_=Depends(require_admin)):
    with WB_LOCK:
        try:
            st = _sync_state(r.year)
            added = 0
            if r.keep_workbook_rows and st["only_wb"]:
                recs = [{k: v for k, v in x.items() if k not in ("row", "date_issue")} for x in st["only_wb"]]
                added = db.import_upload(recs, me_["username"], uuid.uuid4().hex[:12], f"Sync {os.path.basename(st['path'])}")
                st = _sync_state(st["year"])
            items = {s: [] for s in st["book"]}
            order = []
            for x, s in st["led"]:
                f = x["fields"]
                items[s].append(dict(id=x["id"], sheet=s, **{k: f.get(k, "") for k in ("received", "customer", "order", "ticket", "payref",
                                                                                          "method", "amount", "taxable", "deposit_bank")}))
            for s in items: items[s].sort(key=lambda it: (it["received"], st["pos"].get(it["id"], 10 ** 9), it["id"]))
            placed = writer.rebuild_sheets(st["path"], items)
        except writer.WorkbookError as e: raise HTTPException(409, str(e))
    fields = {x["id"]: x["fields"] for x, _ in st["led"]}
    db.link_synced(st["path"], placed, fields)
    unlinked = db.unlink_rows(st["path"], set(placed))
    db.audit(me_["username"], "workbook_synced", file=os.path.basename(st["path"]), year=st["year"], rows=len(placed),
             added_to_ledger=added, removed=len(st["only_wb"]) if not r.keep_workbook_rows else 0)
    return dict(written=len(placed), added_to_ledger=added, year=st["year"], workbook=os.path.basename(st["path"]),
                removed=0 if r.keep_workbook_rows else len(st["only_wb"]), unlinked=unlinked)

@app.get("/api/workbook/download")
def download_current_workbook(me_=Depends(require_perm("export"))):
    """The workbook in use, as it is in Excel (all payee sheets, day blocks, subtotals)."""
    p = workbook_path()
    if not os.path.isfile(p): raise HTTPException(404, "The workbook file isn't there. An admin can pick one in Settings.")
    db.audit(me_["username"], "workbook_downloaded", file=os.path.basename(p))
    return FileResponse(p, filename=os.path.basename(p), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

class LdapTest(BaseModel):
    username: str
    password: str
    settings: dict = {}

@app.post("/api/settings/test-ldap")
def test_ldap(t: LdapTest, _=Depends(require_admin)):
    from . import ldap_auth
    s = db.settings() | {k: v for k, v in t.settings.items() if k.startswith("ldap_")}
    if not s.get("ldap_server"): raise HTTPException(400, "Enter the AD server first.")
    return ldap_auth.test(s, t.username, t.password)

# ================================================================== upload + background extraction
JOBS: "queue.Queue[tuple]" = queue.Queue()

def _worker():
    from . import extract
    while True:
        doc_id, path, filename, received, sheet_hint = JOBS.get()
        try:
            db.update_document(doc_id, status="processing", progress=0.02)
            reg = db.registry()
            def prog(k, n): db.update_document(doc_id, progress=round(0.05 + 0.9 * k / max(n, 1), 3), pages=n)
            data = open(path, "rb").read()
            res = extract.process(data, filename, reg, received or None, progress=prog)
            keep_sensitive = db.get_setting("keep_sensitive_pages") == "1"
            page_ids = {}
            for pg in res["pages"]:
                img = pg["image"]
                if pg["sensitive"] and not keep_sensitive:       # ID / card photos: keep only an unreadable placeholder
                    img = cv2.resize(cv2.resize(img, (24, 32)), (img.shape[1] // 4, img.shape[0] // 4), interpolation=cv2.INTER_NEAREST)
                name = f"{uuid.uuid4().hex}.jpg"
                with open(os.path.join(config.IMAGES, name), "wb") as fh: fh.write(extract.jpeg(img))
                page_ids[pg["index"]] = db.add_page(doc_id, pg["index"], pg["kind"], name, pg["sensitive"])
            for p in res["payments"]:
                f = dict(p["fields"])
                if sheet_hint and not f.get("sheet"):
                    f["sheet"] = sheet_hint; p["warns"].pop("sheet", None)
                if not f.get("received"): f["received"] = received or datetime.date.today().isoformat()
                db.add_payment(doc_id, f, p["warns"], p["sources"], [page_ids[i] for i in p["pages"] if i in page_ids])
            unread = sum(1 for p in res["payments"] if "page" in p["warns"])
            db.update_document(doc_id, status="done", progress=1, pages=len(res["pages"]), received=res.get("received") or received or "",
                               error=(f"{unread} page{'s' if unread > 1 else ''} need typing in" if unread else None),
                               found=len(res["payments"]))
        except Exception as e:
            traceback.print_exc()
            db.update_document(doc_id, status="error", error=f"{type(e).__name__}: {e}"[:400])
        finally:
            try: os.remove(path)
            except OSError: pass
            JOBS.task_done()

threading.Thread(target=_worker, daemon=True).start()

@app.post("/api/documents")
async def upload(files: list[UploadFile] = File(...), received: str = Form(""), sheet: str = Form(""), user=Depends(require_admin)):
    out = []
    os.makedirs(os.path.join(config.DATA, "uploads"), exist_ok=True)
    for f in files:
        name = f.filename or "upload"
        if not name.lower().endswith((".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".heic", ".webp", ".bmp")):
            raise HTTPException(400, f"{name}: upload PDF or image files.")
        path = os.path.join(config.DATA, "uploads", uuid.uuid4().hex + os.path.splitext(name)[1].lower())
        with open(path, "wb") as fh: fh.write(await f.read())
        doc_id = db.add_document(name, received, user["username"])
        JOBS.put((doc_id, path, name, received, sheet))
        out.append(db.get_document(doc_id))
    db.audit(user["username"], "uploaded", files=[o["filename"] for o in out])
    return out

@app.get("/api/documents/{doc_id}")
def document(doc_id: int, _=Depends(require_admin)):
    d = db.get_document(doc_id)
    if not d: raise HTTPException(404)
    return d

def _file_stem(f):
    """A readable download name: 2026-10-09_ACME-BUILDERS_1250.00"""
    import re as _re
    cust = _re.sub(r"[^A-Za-z0-9]+", "-", (f.get("customer") or "payment").upper()).strip("-")[:40] or "PAYMENT"
    return f"{f.get('received') or 'undated'}_{cust}_{f.get('amount') or ''}".rstrip("_")

@app.get("/api/pages/{page_id}")
def page_image(page_id: int, download: int = 0, user=Depends(require_perm("images.view"))):
    p = db.page_file(page_id)
    if not p: raise HTTPException(404)
    if p["sensitive"] and user["role"] != "admin": raise HTTPException(403)
    if download and not can(user, "images.download"): raise HTTPException(403, "Your account can view images but not download them.")
    path = os.path.join(config.IMAGES, p["image"])
    if not os.path.exists(path): raise HTTPException(404)
    headers = {"Cache-Control": "private, max-age=86400"}
    if download:
        with db.conn() as c:
            r = c.execute("""SELECT pay.fields, pp.ord FROM payment_pages pp JOIN payments pay ON pay.id=pp.payment_id WHERE pp.page_id=?""", (page_id,)).fetchone()
        f = json.loads(r["fields"]) if r else {}
        headers["Content-Disposition"] = f'attachment; filename="{_file_stem(f)}_p{(r["ord"] + 1) if r else 1}.jpg"'
        db.audit(user["username"], "image_downloaded", page=page_id)
    return FileResponse(path, media_type="image/jpeg", headers=headers)

@app.get("/api/payments/{pid}/document.pdf")
def payment_pdf(pid: int, user=Depends(require_perm("images.download"))):
    """Every page of a payment (scans + attachments) as one PDF, e.g. to attach the payment in NetSuite."""
    from io import BytesIO
    from PIL import Image
    pay = db.get_payment(pid)
    if not pay: raise HTTPException(404)
    pages = [p for p in db.payment_pages(pid) if user["role"] == "admin" or not p["sensitive"]]
    imgs = []
    for p in pages:
        fp = os.path.join(config.IMAGES, p["image"])
        if os.path.exists(fp): imgs.append(Image.open(fp).convert("RGB"))
    if not imgs: raise HTTPException(404, "This payment has no pages to download.")
    buf = BytesIO()
    imgs[0].save(buf, "PDF", resolution=150, save_all=True, append_images=imgs[1:])
    db.audit(user["username"], "pdf_downloaded", pid, pages=len(imgs))
    return Response(buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{_file_stem(pay["fields"])}.pdf"'})

ATTACH_EXT = (".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp", ".bmp")

def _attachment_images(data, name):
    """PDF pages (150 dpi) or one photo -> PIL images, large photos scaled to 2200 px on the long side."""
    from io import BytesIO
    from PIL import Image, ImageOps
    out = []
    if name.lower().endswith(".pdf"):
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(data)
        if len(doc) > 40: raise HTTPException(400, f"{name}: more than 40 pages. Split it first.")
        for i in range(len(doc)): out.append(doc[i].render(scale=150 / 72).to_pil().convert("RGB"))
    else:
        try: im = ImageOps.exif_transpose(Image.open(BytesIO(data))).convert("RGB")
        except Exception: raise HTTPException(400, f"{name}: couldn't open this image.")
        out.append(im)
    for k, im in enumerate(out):
        if max(im.size) > 2200:
            im.thumbnail((2200, 2200)); out[k] = im
    return out

@app.post("/api/payments/{pid}/attachments")
async def add_attachments(pid: int, files: list[UploadFile] = File(...), user=Depends(require_perm("attachments.edit"))):
    """Extra pages for a payment (e.g. the 3 pages of a card payment, a paper receipt). Stored as-is, not read by OCR."""
    if not db.get_payment(pid): raise HTTPException(404)
    page_ids = []
    for f in files:
        name = os.path.basename(f.filename or "attachment")
        if not name.lower().endswith(ATTACH_EXT): raise HTTPException(400, f"{name}: attach a PDF or an image (JPG, PNG, TIFF). iPhone HEIC photos: export as JPG first.")
        for im in _attachment_images(await f.read(), name):
            fn = f"{uuid.uuid4().hex}.jpg"
            im.save(os.path.join(config.IMAGES, fn), "JPEG", quality=85)
            page_ids.append(db.add_page(None, 0, "attachment", fn, False))
    db.attach_pages(pid, page_ids)
    db.audit(user["username"], "attachment_added", pid, pages=len(page_ids), files=[os.path.basename(f.filename or "") for f in files])
    return db.get_payment(pid)

@app.delete("/api/payments/{pid}/attachments/{page_id}")
def remove_attachment(pid: int, page_id: int, user=Depends(require_perm("attachments.edit"))):
    img = db.remove_attachment(pid, page_id)
    if not img: raise HTTPException(404, "Only pages added as attachments can be removed.")
    try: os.remove(os.path.join(config.IMAGES, img))
    except OSError: pass
    db.audit(user["username"], "attachment_removed", pid, page=page_id)
    return db.get_payment(pid)

# ================================================================== review queue
@app.get("/api/payments")
def payments(status: str = "review,validated", _=Depends(require_admin)):
    return db.list_payments(status.split(","))

class NewPayment(BaseModel):
    fields: dict = {}

@app.post("/api/payments")
def new_payment(p: NewPayment, user=Depends(require_admin)):
    """Manual entry (e.g. cash with no paperwork scanned)."""
    f = dict(p.fields); f.setdefault("method", "cash"); f.setdefault("received", datetime.date.today().isoformat())
    if f.get("amount") not in (None, ""):                      # '$1,000.5' typed by hand -> '1000.50', like edits do
        try: f["amount"] = f"{float(str(f['amount']).replace(',', '').replace('$', '').strip()):.2f}"
        except ValueError: raise HTTPException(400, "Amount must be a number.")
    if f.get("customer"): f.update({k: v for k, v in db.registry().defaults_for(f["customer"], f["method"]).items() if not f.get(k)})
    pid = db.add_payment(None, f, {}, ["manual"], [])
    db.audit(user["username"], "payment_added", pid)
    return db.get_payment(pid)

class Patch(BaseModel):
    fields: dict | None = None
    status: str | None = None

@app.patch("/api/payments/{pid}")
def patch_payment(pid: int, p: Patch, user=Depends(require_admin)):
    if p.status and p.status not in ("review", "validated", "pending"): raise HTTPException(400, "Bad status")
    try: r = db.update_payment(pid, p.fields, p.status)
    except ValueError as e: raise HTTPException(409, str(e))
    if not r: raise HTTPException(404)
    if p.status: db.audit(user["username"], "status_" + p.status, pid)
    return r

@app.delete("/api/payments/{pid}")
def delete_payment(pid: int, user=Depends(require_admin)):
    try: db.delete_payment(pid)
    except ValueError as e: raise HTTPException(409, str(e))
    db.audit(user["username"], "payment_deleted", pid)
    return {"ok": True}

class PageMove(BaseModel):
    page_id: int
    to: int | None = None

@app.post("/api/payments/{pid}/move-page")
def move_page(pid: int, m: PageMove, user=Depends(require_admin)):
    new = db.move_page(m.page_id, pid, m.to)
    return dict(source=db.get_payment(pid), target=db.get_payment(new))

# ================================================================== submit to Excel
class Submit(BaseModel):
    ids: list[int]

REQUIRED = {"customer": "customer", "sheet": "payee sheet", "amount": "amount", "received": "date received", "method": "payment method"}

@app.post("/api/submit")
def submit(s: Submit, user=Depends(require_admin)):
    items, skipped, rows = [], [], {p["id"]: p for p in db.list_payments(["validated"])}
    valid, payees = sheets(), db.registry().payees()
    for pid in s.ids:
        p = rows.get(pid)
        if not p: skipped.append(dict(id=pid, reason="Not validated.")); continue
        f = p["fields"]
        missing = [label for k, label in REQUIRED.items() if not f.get(k)]
        if missing: skipped.append(dict(id=pid, reason="Missing " + ", ".join(missing) + ".")); continue
        target = book_sheet(f["sheet"], valid, payees)
        if not target: skipped.append(dict(id=pid, reason=f"Sheet '{f['sheet']}' is not in the workbook.")); continue
        if (p["warns"].get("duplicate") or "").startswith("Already sent"): skipped.append(dict(id=pid, reason=p["warns"]["duplicate"])); continue
        items.append(dict(id=pid, **{k: f.get(k, "") for k in ("received", "customer", "order", "ticket", "payref", "method",
                                                                 "amount", "taxable", "deposit_bank")}, sheet=target))
    written = {}
    if items:
        path = workbook_path()
        try:
            with WB_LOCK: written = writer.write_payments(path, items)
        except writer.WorkbookError as e:
            raise HTTPException(409, str(e))
        reg = db.registry()
        for it in items:
            sheet, row = written[it["id"]]
            db.mark_submitted(it["id"], path, sheet, row, user["username"])
            reg.learn(rows[it["id"]]["fields"])
        _relocate({(it["sheet"], it["received"]) for it in items})
        db.audit(user["username"], "submitted", ids=list(written), sheets=sorted({v[0] for v in written.values()}))
    per = {}
    for sheet, _ in written.values(): per[sheet] = per.get(sheet, 0) + 1
    return dict(written=len(written), sheets=per, skipped=skipped, workbook=workbook_path())

def _relocate(days, path=None):
    """After a day block is rebuilt, earlier rows of that day moved: refresh the remembered row numbers."""
    import openpyxl
    from .excel.layout import Layout
    path = path or workbook_path()
    try: wb = openpyxl.load_workbook(path, read_only=False)
    except Exception: return
    upd = {}
    for sheet, day in days:
        if sheet not in wb.sheetnames: continue
        lay = Layout(wb[sheet])
        for p in db.submitted_on(sheet, day, path):
            r = writer.find_row(lay, p["fields"], p["row"])
            if r and r != p["row"]: upd[p["id"]] = r
    db.set_rows(upd)

# ================================================================== ledger
@app.get("/api/ledger")
def ledger(user=Depends(require_perm("ledger.view"))):
    rows = db.ledger_rows()
    if user["role"] != "admin":
        for r in rows: r["pages"] = [p for p in r["pages"] if not p["sensitive"]]
    return dict(rows=rows, sheets=sheets(), deposit_banks=[b.strip() for b in (db.get_setting("deposit_banks") or "").split(",") if b.strip()])

# Ledger edits: AR may set Taxable; admins may correct any field of a payment already sent to Excel.
# Every edit takes the same road: saved at once, row queued, written to the workbook now or by the 30 s retry loop.
EXCEL_KEYS = ["sheet", "received", "customer", "order", "ticket", "payref", "method", "amount", "taxable", "deposit_bank"]
EDITABLE = EXCEL_KEYS + ["check_no", "bank", "invoice", "confirmation", "memo"]
WB_OPEN = "The workbook is open in Excel. Your change is saved and will be written automatically once it is closed."

def push_to_excel():
    """Write queued ledger changes. Each payment goes back to the workbook it was written to, so corrections to last
    year's payments land in last year's file even after a new workbook is selected."""
    rows = db.dirty_rows()
    if not rows: return 0, ""
    pick = lambda f: {k: f.get(k, "") for k in EXCEL_KEYS}
    by_book = {}
    for r in rows: by_book.setdefault(r["workbook"] or workbook_path(), []).append(r)
    left, msg = 0, ""
    for path, rs in by_book.items():
        if not os.path.exists(path):
            m = f"Its workbook ({os.path.basename(path)}) isn't there anymore, so Excel can't be updated."
            for r in rs: db.sync_result(r["id"], note=m)
            left, msg = left + len(rs), m; continue
        try: bs, payees = writer.list_sheets(path), db.registry().payees()
        except Exception: bs, payees = [], []
        def mapped(f):
            f = pick(f); f["sheet"] = book_sheet(f.get("sheet"), bs, payees) or f.get("sheet"); return f
        items = [dict(id=r["id"], row=r["row"], op=r["op"], old=mapped(r["excel"]), new=mapped(r["fields"])) for r in rs]
        try:
            with WB_LOCK: res = writer.sync_rows(path, items)
        except writer.WorkbookError as e:
            m = WB_OPEN if "open in Excel" in str(e) else str(e)
            for r in rs: db.sync_result(r["id"], note=m)
            left, msg = left + len(rs), m; continue
        moved = set()
        fields = {r["id"]: r["fields"] for r in rs}
        for pid, out in res.items():
            if out[0] == "removed":
                db.sync_result(pid, removed=True)
            elif out[0] == "ok":
                _, sheet, row, mv = out
                db.sync_result(pid, row=row, sheet=sheet, fields=fields[pid])
                if mv: moved.add((sheet, fields[pid].get("received")))
            else:
                db.sync_result(pid, note=out[1]); left += 1; msg = out[1]
        if moved: _relocate(moved, path)
    return left, msg

def _clean_edit(before, body, user):
    """Validate one ledger edit -> {field: new value} (only what actually changes)."""
    unknown = [k for k in body if k not in EDITABLE]
    if unknown: raise HTTPException(400, "These fields can't be edited: " + ", ".join(unknown))
    if user["role"] != "admin":
        if not can(user, "taxable.edit"): raise HTTPException(403, "Your account can't change payments.")
        if set(body) - {"taxable"}: raise HTTPException(403, "You can only change the Taxable field.")
    old = before["fields"]
    new = {k: ("" if v is None else str(v).strip()) for k, v in body.items()}
    if "amount" in new:
        try: amt = float(new["amount"].replace(",", "").replace("$", ""))
        except ValueError: raise HTTPException(400, "Amount must be a number.")
        if amt == 0: raise HTTPException(400, "Amount can't be zero (use a negative amount for a refund or reimbursement).")
        new["amount"] = f"{amt:.2f}"
    if "received" in new:
        try: datetime.date.fromisoformat(new["received"])
        except ValueError: raise HTTPException(400, "Date received must be a valid date.")
    if "method" in new and new["method"] not in config.METHODS: raise HTTPException(400, "Unknown payment method.")
    if "taxable" in new and new["taxable"] not in ("Yes", "No", ""): raise HTTPException(400, "Taxable must be Yes, No or blank.")
    if "customer" in new:
        new["customer"] = " ".join(new["customer"].upper().split())
        if not new["customer"]: raise HTTPException(400, "Customer can't be empty.")
    if "deposit_bank" in new: new["deposit_bank"] = new["deposit_bank"].upper()
    if "sheet" in new:
        if not new["sheet"]: raise HTTPException(400, "Choose the payee sheet.")
        if (before.get("excel") or {}).get("row") and new["sheet"] not in sheets(): raise HTTPException(400, f"Sheet '{new['sheet']}' is not in the workbook.")
    if "payref" in new and (new.get("method") or old.get("method")) == "check" and "check_no" not in new: new["check_no"] = new["payref"]
    return {k: v for k, v in new.items() if (old.get(k) or "") != v}

@app.patch("/api/ledger/{pid}")
def ledger_patch(pid: int, body: dict, user=Depends(require_user)):
    before = db.get_payment(pid)
    if not before or before["status"] != "submitted": raise HTTPException(404, "Payment not found in the ledger.")
    ch = _clean_edit(before, body, user)
    if not ch: return dict(ok=True, changed=[], excel="unchanged", message="")
    db.set_ledger(pid, **ch)
    db.audit(user["username"], "ledger_edit", pid, **{k: [before["fields"].get(k, ""), v] for k, v in ch.items()})
    if "customer" in ch: db.registry().learn(db.get_payment(pid)["fields"])
    exc = before.get("excel") or {}
    if not exc.get("row") or not any(k in EXCEL_KEYS for k in ch):
        if not exc.get("pending"): db.sync_result(pid)          # nothing for Excel (an earlier queued change still waits)
        return dict(ok=True, changed=sorted(ch), excel="pending" if exc.get("pending") else "skipped", message="")
    left, msg = push_to_excel()
    after = db.get_payment(pid)
    pending = bool((after.get("excel") or {}).get("pending"))
    return dict(ok=True, changed=sorted(ch), excel="pending" if pending else "synced", message=msg if pending else "",
                sheet=(after.get("excel") or {}).get("sheet"), row=(after.get("excel") or {}).get("row"))

class LedgerDelete(BaseModel):
    reason: str = ""

@app.post("/api/ledger/{pid}/delete")
def ledger_delete(pid: int, d: LedgerDelete, user=Depends(require_admin)):
    """Wrongly entered: move it to Deleted (out of the ledger, totals and analytics; its workbook row is emptied).
    Nothing is lost: it can be restored."""
    if not db.ledger_delete(pid, user["username"], d.reason.strip()): raise HTTPException(404, "Payment not found in the ledger.")
    db.audit(user["username"], "ledger_deleted", pid, reason=d.reason.strip())
    left, msg = push_to_excel()
    pending = bool((db.get_payment(pid).get("excel") or {}).get("pending"))
    return dict(ok=True, excel="pending" if pending else "synced", message=msg if pending else "")

@app.get("/api/ledger/deleted")
def ledger_deleted(_=Depends(require_admin)): return db.deleted_rows()

@app.post("/api/ledger/{pid}/restore")
def ledger_restore(pid: int, user=Depends(require_admin)):
    if not db.ledger_restore(pid): raise HTTPException(404, "Payment not found in Deleted.")
    db.audit(user["username"], "ledger_restored", pid)
    left, msg = push_to_excel()
    pending = bool((db.get_payment(pid).get("excel") or {}).get("pending"))
    return dict(ok=True, excel="pending" if pending else "synced", message=msg if pending else "")

class ExportReq(BaseModel):
    ids: list[int] = []

@app.post("/api/ledger/export")
def ledger_export(r: ExportReq, user=Depends(require_perm("export"))):
    """The ledger as an Excel file, plus daily totals per deposit bank. No ids: the whole ledger (the Export Excel button);
    ids: just those rows."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    want = set(r.ids)
    rows = [x for x in db.ledger_rows() if not want or x["id"] in want]
    rows.sort(key=lambda x: (x["fields"].get("received") or "", x["id"]), reverse=True)
    num = lambda v: float(str(v or 0).replace(",", "") or 0)
    meth = {"check": "Check", "cash": "Cash", "card": "Credit card", "ach": "ACH", "wire": "Wire"}
    wb = Workbook(); ws = wb.active; ws.title = "Ledger"
    head = ["Received", "Customer", "Payee", "Method", "Ticket ref", "Payment ref", "Order #", "Amount", "Taxable", "Deposit bank",
            "Check #", "Bank (drawn on)", "Invoice", "Pages", "Excel row"]
    ws.append(head)
    for x in rows:
        f = x["fields"]
        ws.append([f.get("received"), f.get("customer"), x.get("sheet"), meth.get(f.get("method"), f.get("method")), f.get("ticket"), f.get("payref"),
                   f.get("order"), num(f.get("amount")), f.get("taxable"), f.get("deposit_bank"), f.get("check_no"), f.get("bank"), f.get("invoice"),
                   len(x.get("pages") or []), x.get("row")])
    n = len(rows) + 1
    ws.append(["Total", "", "", "", "", "", "", f"=SUBTOTAL(9,H2:H{n})"])
    bold, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="381F12")
    for c in ws[1]: c.font, c.fill = bold, fill
    for c in ws[n + 1]: c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
        for c in row: c.number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJKLMNO", (12, 34, 9, 12, 11, 14, 11, 14, 9, 13, 11, 22, 12, 7, 9)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:O{n}"
    # daily totals per deposit bank (the cashier's view)
    t = wb.create_sheet("Daily bank totals")
    t.append(["Day", "Deposit bank", "Payments", "Cash", "Checks", "Card / ACH / wire", "Total", "Taxable", "Not taxable", "Taxable not set"])
    agg = {}
    for x in rows:
        f = x["fields"]; k = (f.get("received") or "", f.get("deposit_bank") or "No deposit bank")
        a = agg.setdefault(k, [0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0]); v = num(f.get("amount"))
        a[0] += 1; a[1 if f.get("method") == "cash" else 2 if f.get("method") == "check" else 3] += v; a[4] += v
        a[5 if f.get("taxable") == "Yes" else 6 if f.get("taxable") == "No" else 7] += v
    for (day, bank), a in sorted(agg.items(), key=lambda kv: (kv[0][0], kv[0][1]), reverse=True):
        t.append([day, bank, *a])
    for c in t[1]: c.font, c.fill = bold, fill
    for row in t.iter_rows(min_row=2, min_col=4, max_col=10):
        for c in row: c.number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJ", (12, 18, 10, 13, 13, 17, 14, 13, 13, 15)): t.column_dimensions[col].width = w
    t.freeze_panes = "A2"
    buf = BytesIO(); wb.save(buf)
    db.audit(user["username"], "ledger_exported", rows=len(rows))
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="checkflow-ledger-{datetime.date.today().isoformat()}.xlsx"'})

def _sync_loop():
    import time
    while True:
        time.sleep(30)
        try:
            if db.dirty_rows(): push_to_excel()
        except Exception: traceback.print_exc()

threading.Thread(target=_sync_loop, daemon=True).start()

# ================================================================== learned data + import
@app.get("/api/learned")
def learned(_=Depends(require_admin)): return db.learned()

@app.put("/api/learned/{table}")
def learned_put(table: str, row: dict, key: str | None = None, user=Depends(require_admin)):
    if table not in db.TABLES: raise HTTPException(404)
    db.learned_upsert(table, row, key)
    db.audit(user["username"], "learned_" + table, key=key, row=row)
    return db.learned()

@app.delete("/api/learned/{table}")
def learned_del(table: str, key: str, user=Depends(require_admin)):
    if table not in db.TABLES: raise HTTPException(404)
    db.learned_delete(table, key)
    db.audit(user["username"], "learned_delete_" + table, key=key)
    return db.learned()

class ImportReq(BaseModel):
    what: str = "customers"          # customers | history
    sheets: list[str] = []
    since: str = ""

@app.post("/api/import")
def do_import(r: ImportReq, user=Depends(require_admin)):
    path = workbook_path()
    sh = r.sheets or sheets()
    try:
        if r.what == "customers":
            n = db.import_customers(importer.customers(path, sh))
        else:
            recs = [x for x in importer.history(path, sh) if not r.since or (x["received"] or "") >= r.since]
            n = db.import_history(recs, user["username"])
            db.import_customers(importer.customers(path, sh))
    except Exception as e:
        raise HTTPException(400, f"Couldn't read the workbook: {e}")
    db.audit(user["username"], "import_" + r.what, count=n, sheets=sh)
    return dict(imported=n)

# ------------------------------------------------------------------ spreadsheet upload into the ledger
IMPORT_DIR = os.path.join(config.DATA, "imports")

def _import_file(token):
    if not token.isalnum(): raise HTTPException(404)
    for ext in (".xlsx", ".xlsm", ".csv"):
        p = os.path.join(IMPORT_DIR, token + ext)
        if os.path.exists(p): return p
    raise HTTPException(404, "That upload expired. Choose the file again.")

def _read_upload(path, filename, date_policy="fix", parsed=None):
    """Read (or reuse the preview's parsed copy), apply the date policy, then mark rows already in the ledger."""
    from .excel import ledger_import
    if parsed is None:
        try: res = ledger_import.read_file(path, filename, sheets(), db.registry())
        except ValueError as e: raise HTTPException(400, str(e))
    else:
        res = parsed
    res["flagged"] = [dict(source=r.get("source", ""), sheet=r["sheet"], row=r["row"], customer=r["customer"], amount=r["amount"], **r["date_issue"])
                      for r in res["records"] if r.get("date_issue")]
    res["records"], res["fixed"], res["dropped"] = ledger_import.apply_date_policy(res["records"], date_policy)
    ledger_import.mark_duplicates(res["records"], db.ledger_fingerprints())
    return res

@app.post("/api/import/excel")
async def import_excel_preview(file: UploadFile = File(...), user=Depends(require_admin)):
    """Step 1: read the spreadsheet and show what would be added. Nothing is written yet."""
    from .excel import ledger_import
    name = os.path.basename(file.filename or "upload.xlsx")
    ext = os.path.splitext(name.lower())[1]
    if ext not in (".xlsx", ".xlsm", ".csv"): raise HTTPException(400, "Choose an .xlsx, .xlsm or .csv file. (Old .xls files: open in Excel and Save As .xlsx.)")
    os.makedirs(IMPORT_DIR, exist_ok=True)
    for f in os.listdir(IMPORT_DIR):                 # uploads that were never committed
        fp = os.path.join(IMPORT_DIR, f)
        if os.path.getmtime(fp) < datetime.datetime.now().timestamp() - 86400: os.remove(fp)
    token = uuid.uuid4().hex
    path = os.path.join(IMPORT_DIR, token + ext)
    with open(path, "wb") as fh: fh.write(await file.read())
    with open(os.path.join(IMPORT_DIR, token + ".name"), "w", encoding="utf-8") as fh: fh.write(name)
    from .excel import ledger_import as li
    try: parsed = li.read_file(path, name, sheets(), db.registry())
    except ValueError as e:
        os.remove(path); raise HTTPException(400, str(e))
    with open(os.path.join(IMPORT_DIR, token + ".json"), "w", encoding="utf-8") as fh: json.dump(parsed, fh)
    res = _read_upload(path, name, parsed=json.loads(json.dumps(parsed)))
    recs = res["records"]
    new = [r for r in recs if not r["duplicate"]]
    return dict(token=token, filename=name, formats=res["sheets"], warnings=res["warnings"], skipped=res["skipped"],
                flagged_dates=res["flagged"], fixable=sum(1 for f in res["flagged"] if f["fix"]),
                rows=len(recs), new=len(new), duplicates=len(recs) - len(new), total_new=round(sum(float(r["amount"]) for r in new), 2),
                first=min(r["received"] for r in recs), last=max(r["received"] for r in recs),
                per_sheet=ledger_import.summary(recs), needs_sheet=any(not r["sheet"] for r in recs), known_sheets=sheets(),
                sample=[{k: r[k] for k in ("sheet", "received", "customer", "method", "payref", "ticket", "amount", "taxable", "duplicate", "source")}
                        for r in (new[:12] or recs[:12])])

class ImportCommit(BaseModel):
    default_sheet: str = ""
    date_policy: str = "fix"            # fix | skip | keep
    include_duplicates: bool = False
    sheets: list[str] = []
    also_customers: bool = True

@app.post("/api/import/excel/{token}")
def import_excel_commit(token: str, r: ImportCommit, user=Depends(require_admin)):
    """Step 2: add the previewed rows to the ledger as one batch (Undo removes the batch)."""
    path = _import_file(token)
    side = os.path.join(IMPORT_DIR, token + ".name")
    name = open(side, encoding="utf-8").read().strip() if os.path.exists(side) else os.path.basename(path)
    if r.date_policy not in ("fix", "skip", "keep"): raise HTTPException(400, "date_policy must be fix, skip or keep.")
    cache = os.path.join(IMPORT_DIR, token + ".json")
    parsed = json.load(open(cache, encoding="utf-8")) if os.path.exists(cache) else None
    res = _read_upload(path, path, r.date_policy, parsed)
    recs = [x for x in res["records"] if r.include_duplicates or not x["duplicate"]]
    for x in recs:
        if not x["sheet"]: x["sheet"] = r.default_sheet.strip().upper()
    if any(not x["sheet"] for x in recs): raise HTTPException(400, "Some rows don't say which payee they belong to. Pick a payee for them.")
    if r.sheets: recs = [x for x in recs if x["sheet"] in r.sheets]
    batch = token[:16]
    filename = name
    n = db.import_upload(recs, user["username"], batch, filename)
    if r.also_customers and recs:
        names = {}
        for x in recs: names.setdefault(x["customer"], set()).add(x["sheet"])
        db.import_customers({k: sorted(v) for k, v in names.items() if len(k) >= 2})
    for f in (path, side, cache):
        try: os.remove(f)
        except OSError: pass
    db.audit(user["username"], "import_excel", count=n, batch=batch, file=filename)
    return dict(imported=n, batch=batch, skipped_duplicates=sum(1 for x in res["records"] if x["duplicate"]) if not r.include_duplicates else 0)

@app.get("/api/import/batches")
def import_batches(_=Depends(require_admin)): return db.import_batches()

@app.delete("/api/import/batches/{batch}")
def import_undo(batch: str, user=Depends(require_admin)):
    n = db.undo_import(batch)
    if not n: raise HTTPException(404, "Import not found.")
    db.audit(user["username"], "import_undone", count=n, batch=batch)
    return dict(removed=n, batches=db.import_batches())

# ================================================================== analytics
@app.get("/api/analytics")
def analytics_view(start: str = "", end: str = "", grain: str = "month", by: str = "sheet", sheet_list: str = Query("", alias="sheets"), customer: str = "",
                   methods: str = "", banks: str = "", horizon: int = 3, _=Depends(require_admin)):
    from . import analytics
    split = lambda v: [x for x in v.split(",") if x]
    try:
        return analytics.analytics(start or None, end or None, grain, by, split(sheet_list), customer, split(methods), split(banks), horizon,
                                   [b.strip().upper() for b in (db.get_setting("deposit_banks") or "").split(",") if b.strip()])
    except ValueError as e:
        raise HTTPException(400, f"Bad filter: {e}")

@app.get("/api/audit")
def audit_log(limit: int = 200, _=Depends(require_admin)):
    with db.conn() as c:
        return [dict(r) for r in c.execute("SELECT * FROM audit ORDER BY id DESC LIMIT ?", (min(limit, 2000),))]

@app.get("/api/meta")
def meta(_=Depends(require_user)):  # noqa: shared lists (sheets, banks) every role needs for the ledger
    return dict(methods=config.METHODS, sheets=sheets(),
                deposit_banks=[b.strip() for b in (db.get_setting("deposit_banks") or "").split(",") if b.strip()],
                payees=[dict(name=n, sheet=s) for n, s, _ in db.registry().payees()],
                customers=db.registry().customers())

# ================================================================== UI
if os.path.isdir(os.path.join(config.STATIC, "assets")):
    app.mount("/assets", StaticFiles(directory=os.path.join(config.STATIC, "assets")), name="assets")

@app.get("/{path:path}")
def spa(path: str):
    if path.startswith("api/"): raise HTTPException(404)
    f = os.path.join(config.STATIC, path)
    if path and os.path.isfile(f): return FileResponse(f)
    idx = os.path.join(config.STATIC, "index.html")
    if os.path.exists(idx): return FileResponse(idx)
    return JSONResponse({"error": "UI not built. Run: cd frontend && npm install && npm run build"}, status_code=500)
